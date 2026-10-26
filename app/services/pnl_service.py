"""
費用登錄與損益表 (第四輪)：

損益表 (期間內，可依分店)：
  營業額 (含稅、扣除部分退款) − 銷項稅額 = 營業收入淨額
  − 銷貨成本 (配方原料成本或餐點單位成本，見營運報表) = 營業毛利
  − 營業費用：人事成本 (打卡工時 × 時薪) + 金流 / 外送平台手續費 (依收款方式與設定費率估算) + 登錄的各項費用
  = 營業淨利 (淨利率)
並與上一期比較、列出近 6 個月趨勢。原料成本變動對餐點毛利的影響見 cost_impact()。
"""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

from app.errors import AppError
from app.extensions import db
from app.models.enums import PaymentMethod
from app.models.expense import EXPENSE_CATEGORIES, CostHistory, Expense
from app.models.inventory import Ingredient, RecipeLine
from app.models.menu import MenuItem
from app.models.order import Order
from app.models.payment import PaymentKind, PaymentRecord
from app.services import report_service, settings_service
from app.sync import runtime as sync_runtime
from app.utils import neutralize_formulas, safe_int

FEE_GROUPS = {
    'card_fee_percent': (PaymentMethod.CREDIT_CARD, PaymentMethod.APPLE_PAY),
    'mobile_fee_percent': (PaymentMethod.LINE_PAY, PaymentMethod.JKO_PAY),
    'online_fee_percent': (PaymentMethod.ATM, PaymentMethod.CVS),
    'platform_fee_percent': (PaymentMethod.PLATFORM,),
}
FEE_LABELS = {'card_fee_percent': '刷卡手續費', 'mobile_fee_percent': '行動支付手續費', 'online_fee_percent': '線上金流手續費',
              'platform_fee_percent': '外送平台抽成'}


# ---------------------------------------------------------------------------
# 費用
# ---------------------------------------------------------------------------
MIN_YEAR, MAX_YEAR = 2000, 2100   # 查詢範圍 (避免極端年份的日期運算溢位)


def month_range(month):
    """'2026-10' → (date(2026,10,1), date(2026,10,31))；格式錯誤或超出範圍時為本月"""
    try:
        y, m = (int(x) for x in str(month).split('-')[:2])
        if not MIN_YEAR <= y <= MAX_YEAR:
            raise ValueError(month)
        start = date(y, m, 1)
    except (ValueError, TypeError):
        today = date.today()
        start = today.replace(day=1)
    end = (start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    return start, end


def save_expense(exp, form, staff=''):
    name = str(form.get('name') or '').strip()[:80]
    if not name:
        raise AppError('請輸入費用名稱')
    category = form.get('category')
    if category not in EXPENSE_CATEGORIES:
        raise AppError('請選擇費用類別')
    amount = safe_int(form.get('amount'), -1)
    if not 0 < amount <= 100_000_000:
        raise AppError('請輸入正確的金額')
    try:
        spent_on = date.fromisoformat(str(form.get('spent_on') or ''))
    except ValueError:
        raise AppError('請輸入日期')
    exp.name, exp.category, exp.amount, exp.spent_on = name, category, amount, spent_on
    exp.vendor = str(form.get('vendor') or '').strip()[:60]
    exp.invoice_no = str(form.get('invoice_no') or '').strip().upper()[:20]
    exp.is_fixed = bool(form.get('is_fixed'))
    exp.note = str(form.get('note') or '').strip()[:200]
    if exp.id is None:
        exp.created_by = staff
        exp.node_uuid = sync_runtime.local_node_uuid()
    return exp


def expenses(tenant_id, start, end, node_uuid=None):
    q = Expense.query.filter(Expense.tenant_id == tenant_id, Expense.is_deleted.is_(False), Expense.spent_on >= start,
                             Expense.spent_on <= end)
    if node_uuid:
        q = q.filter(Expense.node_uuid == node_uuid)
    return q.order_by(Expense.spent_on, Expense.id).all()


def _own(rows):
    """本機登錄的費用 (線上端的費用 node_uuid 為空；分店上傳到線上端的不算線上端自己的)"""
    local = sync_runtime.local_node_uuid()
    return [e for e in rows if (e.node_uuid or None) == (local or None)]


def copy_fixed(tenant_id, month, staff=''):
    """
    把上個月勾選「每月固定」的費用複製到 month (同名稱已存在的略過)；回傳新增筆數。由呼叫端 commit。
    只複製本機登錄的費用：線上端不會把分店上傳的房租再複製一份 (分店自己複製後會上傳，否則會重複計算)。
    """
    start, _end = month_range(month)
    prev_start, prev_end = month_range(f'{(start - timedelta(days=1)):%Y-%m}')
    existing = {(e.name, e.category) for e in _own(expenses(tenant_id, start, _end))}
    added = 0
    for e in _own(expenses(tenant_id, prev_start, prev_end)):
        if not e.is_fixed or (e.name, e.category) in existing:
            continue
        day = min(e.spent_on.day, _end.day)
        db.session.add(Expense(tenant_id=tenant_id, spent_on=start.replace(day=day), category=e.category, name=e.name,
                               amount=e.amount, vendor=e.vendor, is_fixed=True, note=e.note, created_by=staff,
                               node_uuid=sync_runtime.local_node_uuid()))
        added += 1
    return added


# ---------------------------------------------------------------------------
# 損益表
# ---------------------------------------------------------------------------
def _fees(tenant_id, t0, t1, node_uuid):
    cfg = settings_service.get('pnl', tenant_id)
    q = db.session.query(PaymentRecord.method, db.func.sum(PaymentRecord.amount)).filter(
        PaymentRecord.tenant_id == tenant_id, PaymentRecord.created_at >= t0, PaymentRecord.created_at < t1,
        PaymentRecord.amount > 0,
        PaymentRecord.kind.in_((PaymentKind.SALE, PaymentKind.SETTLE, PaymentKind.TOPUP, PaymentKind.GIFT_SALE)))
    if node_uuid:
        q = q.join(Order, Order.id == PaymentRecord.order_id).filter(Order.node_uuid == node_uuid)
    by_method = dict(q.group_by(PaymentRecord.method).all())
    out = {}
    for key, methods in FEE_GROUPS.items():
        base = sum(int(by_method.get(m) or 0) for m in methods)
        out[key] = round(base * float(cfg.get(key) or 0) / 100)
    return out


def statement(tenant_id, start, end, node_uuid=None):
    r = report_service.build_report(tenant_id, start, end, node_uuid)
    rate = int(settings_service.get('accounting', tenant_id).get('tax_rate') or 0)
    t0 = datetime.combine(start, datetime.min.time())
    t1 = datetime.combine(end + timedelta(days=1), datetime.min.time())
    gross = int(r['net_after_refunds'])
    revenue = round(gross / (1 + rate / 100)) if rate else gross
    tax = gross - revenue
    cogs = int(r['cogs'])
    gross_profit = revenue - cogs
    fees = _fees(tenant_id, t0, t1, node_uuid)
    by_cat = defaultdict(int)
    for e in expenses(tenant_id, start, end, node_uuid):
        by_cat[e.category] += e.amount
    labor = int(r.get('labor_cost') or 0)
    opex = labor + sum(fees.values()) + sum(by_cat.values())
    net = gross_profit - opex
    return {
        'start': start, 'end': end, 'orders': r['order_count'], 'gross_sales': gross, 'tax': tax, 'revenue': revenue,
        'cogs': cogs, 'gross_profit': gross_profit, 'gross_margin': _pct(gross_profit, revenue),
        'labor': labor, 'fees': fees, 'fee_total': sum(fees.values()),
        'expenses': {k: by_cat.get(k, 0) for k in EXPENSE_CATEGORIES}, 'expense_total': sum(by_cat.values()),
        'opex': opex, 'net': net, 'net_margin': _pct(net, revenue),
        'cogs_ratio': _pct(cogs, revenue), 'labor_ratio': _pct(labor, revenue),
        'labor_missing_wage': r.get('labor_missing_wage') or [],
        # 打卡紀錄與員工時薪只在各門市主機：線上端或依分店檢視時不含人事成本 (淨利會偏高)
        'labor_available': not node_uuid and not sync_runtime.is_hub(),
    }


def _pct(a, b):
    return round(a / b * 100, 1) if b else 0.0


def previous_period(start, end):
    """上一期：整月 (可跨多個月) → 往前同樣月數；其他 → 往前同樣天數"""
    days = (end - start).days + 1
    if start.day == 1 and (end + timedelta(days=1)).day == 1:
        months = (end.year - start.year) * 12 + end.month - start.month + 1
        index = start.year * 12 + start.month - 1 - months
        prev_start = date(index // 12, index % 12 + 1, 1)
        return prev_start, start - timedelta(days=1)
    return start - timedelta(days=days), start - timedelta(days=1)


def trend(tenant_id, months=6, node_uuid=None, today=None):
    today = today or date.today()
    out = []
    first = today.replace(day=1)
    for i in range(months - 1, -1, -1):
        y, m = first.year, first.month - i
        while m <= 0:
            y, m = y - 1, m + 12
        start, end = month_range(f'{y}-{m:02d}')
        s = statement(tenant_id, start, min(end, today), node_uuid)
        out.append({'month': f'{y}-{m:02d}', 'revenue': s['revenue'], 'gross_profit': s['gross_profit'],
                    'opex': s['opex'], 'net': s['net'], 'net_margin': s['net_margin']})
    return out


STATEMENT_ROWS = (
    ('營業額 (含稅)', 'gross_sales'), ('銷項稅額', 'tax'), ('營業收入淨額', 'revenue'), ('銷貨成本', 'cogs'),
    ('營業毛利', 'gross_profit'), ('人事成本', 'labor'), ('手續費', 'fee_total'), ('其他營業費用', 'expense_total'),
    ('營業費用合計', 'opex'), ('營業淨利', 'net'),
)


def to_excel(cur, prev, months, shop_name=''):
    wb = Workbook()
    ws = wb.active
    ws.title = '損益表'
    ws.append(['店家', shop_name])
    ws.append(['本期', f"{cur['start']} ~ {cur['end']}", '上期', f"{prev['start']} ~ {prev['end']}"])
    ws.append([])
    ws.append(['項目', '本期', '上期', '增減'])
    for c in ws[4]:
        c.font = Font(bold=True)
    for label, key in STATEMENT_ROWS:
        ws.append([label, cur[key], prev[key], cur[key] - prev[key]])
        if key == 'fee_total':
            for fkey, flabel in FEE_LABELS.items():
                ws.append([f'  {flabel}', cur['fees'][fkey], prev['fees'][fkey], cur['fees'][fkey] - prev['fees'][fkey]])
        if key == 'expense_total':
            for ckey, clabel in EXPENSE_CATEGORIES.items():
                ws.append([f'  {clabel}', cur['expenses'][ckey], prev['expenses'][ckey],
                           cur['expenses'][ckey] - prev['expenses'][ckey]])
    ws.append(['淨利率 %', cur['net_margin'], prev['net_margin'], round(cur['net_margin'] - prev['net_margin'], 1)])
    s2 = wb.create_sheet('近六個月')
    s2.append(['月份', '營業收入淨額', '營業毛利', '營業費用', '營業淨利', '淨利率 %'])
    for m in months:
        s2.append([m['month'], m['revenue'], m['gross_profit'], m['opex'], m['net'], m['net_margin']])
    neutralize_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# 成本 / 售價歷史
# ---------------------------------------------------------------------------
def ingredient_cost_at(tenant_id, when):
    """{ingredient_id: when 當時的單位成本}：when 之後第一筆變動的舊值；之後沒有變動則為目前成本"""
    current = {i.id: float(i.unit_cost or 0) for i in Ingredient.query.filter_by(tenant_id=tenant_id).all()}
    later = CostHistory.query.filter(CostHistory.tenant_id == tenant_id, CostHistory.kind == 'ingredient',
                                     CostHistory.field == 'unit_cost', CostHistory.changed_at > when) \
        .order_by(CostHistory.changed_at.desc(), CostHistory.id.desc()).all()
    out = dict(current)
    for h in later:   # 由新到舊覆蓋，最後留下的是 when 之後第一筆變動的舊值
        out[h.ref_id] = h.old_value
    return out


def cost_impact(tenant_id, days=30, now=None):
    """
    與 days 天前相比，每份餐點的原料成本變化與毛利率變化 (有配方的餐點)，依毛利影響排序。
    回傳 [{id, name, price, cost_then, cost_now, margin_then, margin_now, delta}]
    """
    now = now or datetime.now()
    then = now - timedelta(days=days)
    cost_then, cost_now = ingredient_cost_at(tenant_id, then), ingredient_cost_at(tenant_id, now)
    recipes = defaultdict(list)
    for rl in RecipeLine.query.filter_by(tenant_id=tenant_id).all():
        recipes[rl.menu_item_id].append(rl)
    out = []
    for item in MenuItem.query.filter(MenuItem.tenant_id == tenant_id, MenuItem.id.in_(list(recipes) or [0])).all():
        a = sum((rl.quantity or 0) * cost_then.get(rl.ingredient_id, 0) for rl in recipes[item.id])
        b = sum((rl.quantity or 0) * cost_now.get(rl.ingredient_id, 0) for rl in recipes[item.id])
        price = item.price or 0
        if abs(b - a) < 0.005:
            continue
        out.append({'id': item.id, 'name': item.name, 'price': price, 'cost_then': round(a, 1), 'cost_now': round(b, 1),
                    'delta': round(b - a, 1), 'margin_then': _pct(price - a, price), 'margin_now': _pct(price - b, price)})
    out.sort(key=lambda r: -r['delta'])
    return out


def history(tenant_id, kind=None, days=90, limit=300):
    q = CostHistory.query.filter(CostHistory.tenant_id == tenant_id,
                                 CostHistory.changed_at >= datetime.now() - timedelta(days=days))
    if kind:
        q = q.filter(CostHistory.kind == kind)
    return q.order_by(CostHistory.changed_at.desc(), CostHistory.id.desc()).limit(limit).all()
