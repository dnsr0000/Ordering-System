"""
進階營運報表：營業額、毛利 (配方成本或餐點單位成本)、時段熱度、付款方式、管道、分類與品項排行、
分店比較 (線上端)、收銀員、促銷成效；可匯出 Excel 並定期寄送。
"""
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

from app.extensions import db
from app.models.enums import PAYMENT_METHOD_LABELS, OrderStatus
from app.models.inventory import RecipeLine
from app.models.menu import MenuItem
from app.models.order import Order
from app.models.payment import PaymentRecord, Refund
from app.services import settings_service
from app.utils import neutralize_formulas

REVENUE_STATUSES = ('', 'settled', 'paid')
WEEKDAYS = '一二三四五六日'
SOURCE_LABELS = {'Kiosk': '點餐機', 'POS': 'POS', 'Online': '線上點餐', 'Table': '桌邊掃碼', 'UberEats': 'Uber Eats',
                 'foodpanda': 'foodpanda', 'Delivery': '外送 (手動)'}


def _unit_costs(tenant_id):
    """{menu_item_id: 每份成本}：有配方用原料成本，否則用餐點設定的單位成本"""
    costs = {m.id: float(m.unit_cost or 0) for m in MenuItem.query.filter_by(tenant_id=tenant_id).all()}
    recipe = defaultdict(float)
    for rl in RecipeLine.query.filter_by(tenant_id=tenant_id).all():
        if rl.ingredient is not None:
            recipe[rl.menu_item_id] += (rl.quantity or 0) * (rl.ingredient.unit_cost or 0)
    costs.update(recipe)
    return costs


def build_report(tenant_id, start, end, node_uuid=None):
    """start / end 為 date (含兩端)"""
    t0 = datetime.combine(start, datetime.min.time())
    t1 = datetime.combine(end + timedelta(days=1), datetime.min.time())
    q = Order.query.filter(Order.tenant_id == tenant_id, Order.created_at >= t0, Order.created_at < t1)
    if node_uuid:
        q = q.filter(Order.node_uuid == node_uuid)
    all_orders = q.all()
    orders = [o for o in all_orders if o.status != OrderStatus.CANCELLED and (o.payment_status or '') in REVENUE_STATUSES]
    costs = _unit_costs(tenant_id)
    menu = {m.id: m for m in MenuItem.query.filter_by(tenant_id=tenant_id).all()}

    gross = sum(o.subtotal for o in orders)
    net = sum(o.total_price or 0 for o in orders)
    refunds_q = Refund.query.filter(Refund.tenant_id == tenant_id, Refund.created_at >= t0, Refund.created_at < t1)
    if node_uuid:   # 依分店：只算該分店訂單的退款
        refunds_q = refunds_q.join(Order, Order.id == Refund.order_id).filter(Order.node_uuid == node_uuid)
    partial_refunds = sum(r.amount for r in refunds_q.all() if r.kind == 'partial')
    cogs = 0.0
    items = {}
    categories = defaultdict(lambda: {'qty': 0, 'revenue': 0, 'cost': 0.0})
    for o in orders:
        for it in o.items:
            cost = costs.get(it.menu_item_id, 0.0) * (it.quantity or 0)
            cogs += cost
            key = it.item_name
            row = items.setdefault(key, {'name': key, 'qty': 0, 'revenue': 0, 'cost': 0.0})
            row['qty'] += it.quantity or 0
            row['revenue'] += (it.price or 0) * (it.quantity or 0)
            row['cost'] += cost
            cat = (menu[it.menu_item_id].category if it.menu_item_id in menu else '其他') or '未分類'
            categories[cat]['qty'] += it.quantity or 0
            categories[cat]['revenue'] += (it.price or 0) * (it.quantity or 0)
            categories[cat]['cost'] += cost
    for row in items.values():
        row['margin'] = row['revenue'] - row['cost']
        row['margin_rate'] = round(row['margin'] / row['revenue'] * 100, 1) if row['revenue'] else 0
    by_day = defaultdict(lambda: {'count': 0, 'amount': 0})
    heat = [[0] * 24 for _ in range(7)]
    heat_amount = [[0] * 24 for _ in range(7)]
    by_source = defaultdict(lambda: {'count': 0, 'amount': 0})
    by_cashier = defaultdict(lambda: {'count': 0, 'amount': 0})
    by_node = defaultdict(lambda: {'count': 0, 'amount': 0})
    promotions = defaultdict(lambda: {'count': 0, 'amount': 0})
    for o in orders:
        d = o.created_at.date().isoformat()
        by_day[d]['count'] += 1
        by_day[d]['amount'] += o.total_price or 0
        wd, hr = o.created_at.weekday(), o.created_at.hour
        heat[wd][hr] += 1
        heat_amount[wd][hr] += o.total_price or 0
        src = SOURCE_LABELS.get(o.order_source or 'Kiosk', o.order_source or 'Kiosk')
        by_source[src]['count'] += 1
        by_source[src]['amount'] += o.total_price or 0
        if o.cashier_username:
            by_cashier[o.cashier_username]['count'] += 1
            by_cashier[o.cashier_username]['amount'] += o.total_price or 0
        by_node[o.node_uuid or '本機']['count'] += 1
        by_node[o.node_uuid or '本機']['amount'] += o.total_price or 0
        if o.promotion_discount and o.promotion_names:
            for name in o.promotion_names.split('、'):
                promotions[name]['count'] += 1
                promotions[name]['amount'] += o.promotion_discount // max(1, len(o.promotion_names.split('、')))
    pay_q = PaymentRecord.query.filter(PaymentRecord.tenant_id == tenant_id, PaymentRecord.created_at >= t0,
                                       PaymentRecord.created_at < t1)
    by_method = defaultdict(lambda: {'in': 0, 'out': 0, 'net': 0})
    for p in pay_q.all():
        if p.kind in ('pay_in', 'pay_out'):
            continue
        label = PAYMENT_METHOD_LABELS.get(p.method, p.method)
        if p.amount >= 0:
            by_method[label]['in'] += p.amount
        else:
            by_method[label]['out'] += -p.amount
        by_method[label]['net'] += p.amount
    days = (end - start).days + 1
    report = {
        'start': start.isoformat(), 'end': end.isoformat(), 'days': days,
        'order_count': len(orders), 'cancelled_count': sum(1 for o in all_orders if o.status == OrderStatus.CANCELLED),
        'gross': gross, 'discounts': sum(o.discount_amount or 0 for o in orders),
        'service_charge': sum(o.service_charge or 0 for o in orders), 'net': net,
        'partial_refunds': partial_refunds, 'net_after_refunds': net - partial_refunds,
        'cogs': round(cogs), 'gross_margin': round(net - partial_refunds - cogs),
        'margin_rate': round((net - partial_refunds - cogs) / net * 100, 1) if net else 0,
        'avg_ticket': round(net / len(orders)) if orders else 0,
        'per_day': round(net / days) if days else 0,
        'by_day': [dict(day=k, **v) for k, v in sorted(by_day.items())],
        'heat': heat, 'heat_amount': heat_amount, 'heat_max': max((max(r) for r in heat), default=0),
        'by_source': dict(by_source), 'by_method': dict(by_method), 'by_cashier': dict(by_cashier),
        'by_node': dict(by_node), 'promotions': dict(promotions),
        'categories': {k: dict(v, cost=round(v['cost']), margin=round(v['revenue'] - v['cost'])) for k, v in categories.items()},
        'items': sorted(items.values(), key=lambda r: -r['revenue']),
    }
    report['day_max'] = max((d['amount'] for d in report['by_day']), default=0)
    report.update(labor_cost(tenant_id, t0, t1, report['net_after_refunds']) if not node_uuid else
                  {'labor_hours': 0, 'labor_cost': 0, 'labor_ratio': 0, 'labor_by_staff': [], 'labor_missing_wage': []})
    return report


def labor_cost(tenant_id, t0, t1, revenue):
    """
    人事成本：期間內的打卡工時 × 員工時薪 (員工帳號頁設定)，與營業額相比的人事成本率。
    尚未下班的紀錄以現在時間計算；沒有設定時薪的員工列在 labor_missing_wage 提醒補填。
    """
    from app.models.staff import TimeClock
    now = datetime.now()
    by_staff = {}
    for r in TimeClock.query.filter(TimeClock.tenant_id == tenant_id, TimeClock.clock_in >= t0, TimeClock.clock_in < t1) \
            .execution_options(skip_tenant_filter=True).all():
        end = r.clock_out or min(now, t1)
        hours = max(0.0, (end - r.clock_in).total_seconds() / 3600)
        row = by_staff.setdefault(r.staff_account_id, {'name': r.staff.username if r.staff else str(r.staff_account_id),
                                                       'wage': (r.staff.hourly_wage or 0) if r.staff else 0, 'hours': 0.0})
        row['hours'] += hours
    rows = [dict(v, hours=round(v['hours'], 2), cost=round(v['hours'] * v['wage'])) for v in by_staff.values()]
    rows.sort(key=lambda r: -r['cost'])
    total = sum(r['cost'] for r in rows)
    return {'labor_hours': round(sum(r['hours'] for r in rows), 2), 'labor_cost': total,
            'labor_ratio': round(total / revenue * 100, 1) if revenue else 0, 'labor_by_staff': rows,
            'labor_missing_wage': [r['name'] for r in rows if not r['wage'] and r['hours'] > 0]}


def report_to_excel(report, shop_name=''):
    wb = Workbook()
    ws = wb.active
    ws.title = '總覽'
    rows = [('店家', shop_name), ('期間', f"{report['start']} ~ {report['end']}"), ('訂單數', report['order_count']),
            ('取消單數', report['cancelled_count']), ('商品小計', report['gross']), ('折扣', report['discounts']),
            ('服務費', report['service_charge']), ('營業額', report['net']), ('部分退款', report['partial_refunds']),
            ('退款後營業額', report['net_after_refunds']), ('估計成本', report['cogs']), ('毛利', report['gross_margin']),
            ('毛利率 %', report['margin_rate']), ('客單價', report['avg_ticket']), ('日均營業額', report['per_day'])]
    for r in rows:
        ws.append(r)
    for cell in ws['A']:
        cell.font = Font(bold=True)

    def sheet(title, header, data):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font = Font(bold=True)
        for row in data:
            s.append(row)

    sheet('每日', ['日期', '訂單數', '營業額'], [(d['day'], d['count'], d['amount']) for d in report['by_day']])
    sheet('品項', ['品項', '數量', '營收', '成本', '毛利', '毛利率%'],
          [(i['name'], i['qty'], i['revenue'], round(i['cost']), round(i['margin']), i['margin_rate']) for i in report['items']])
    sheet('分類', ['分類', '數量', '營收', '成本', '毛利'],
          [(k, v['qty'], v['revenue'], v['cost'], v['margin']) for k, v in report['categories'].items()])
    sheet('付款方式', ['方式', '收款', '退款', '淨額'], [(k, v['in'], v['out'], v['net']) for k, v in report['by_method'].items()])
    sheet('點餐管道', ['管道', '訂單數', '營業額'], [(k, v['count'], v['amount']) for k, v in report['by_source'].items()])
    sheet('時段熱度', ['星期'] + [f'{h}時' for h in range(24)],
          [[f'週{WEEKDAYS[i]}'] + report['heat'][i] for i in range(7)])
    sheet('收銀員', ['帳號', '訂單數', '營業額'], [(k, v['count'], v['amount']) for k, v in report['by_cashier'].items()])
    if len(report['by_node']) > 1:
        sheet('分店', ['分店', '訂單數', '營業額'], [(k, v['count'], v['amount']) for k, v in report['by_node'].items()])
    neutralize_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def period_for(frequency, today=None):
    today = today or date.today()
    if frequency == 'weekly':
        end = today - timedelta(days=today.weekday() + 1)      # 上週日
        return end - timedelta(days=6), end
    y = today - timedelta(days=1)
    return y, y


def send_scheduled_report(tenant_id, today=None, force=False):
    """依「定期寄送報表」設定寄出 (每週頻率只在週一寄)；回傳寄出的收件人數"""
    from app.models.tenant import Tenant
    from app.services import mail_service
    cfg = settings_service.get('report_email', tenant_id)
    if not cfg.get('enabled') and not force:
        return 0
    today = today or date.today()
    if cfg.get('frequency') == 'weekly' and today.weekday() != 0 and not force:
        return 0
    start, end = period_for(cfg.get('frequency') or 'daily', today)
    shop = db.session.get(Tenant, tenant_id)
    report = build_report(tenant_id, start, end)
    body = (f"{shop.name if shop else ''} 營運報表 {start} ~ {end}\n\n"
            f"營業額：${report['net']:,}\n訂單數：{report['order_count']}\n客單價：${report['avg_ticket']:,}\n"
            f"毛利 (估計)：${report['gross_margin']:,} ({report['margin_rate']}%)\n\n詳細內容請見附件 Excel。")
    xlsx = report_to_excel(report, shop.name if shop else '').getvalue()
    return mail_service.send_mail(tenant_id, cfg.get('recipients') or '', f'營運報表 {start} ~ {end}', body,
                                  attachments=[(f'report_{start}_{end}.xlsx', xlsx,
                                                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')])
