"""營運報表 (毛利、時段、付款方式、Excel、定期寄送)、員工打卡與班表、系統監控 / 自動備份 / 排程、
出餐追蹤頁、PWA、菜單進階設定與多語系"""
import io
import zipfile
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.menu import MenuItem
from app.models.staff import TimeClock
from app.services import mail_service, maintenance_service, report_service
from tests.support import assert_flash, checkout, default_tenant_id, guest_client, line, login_staff, make_item

POS = '/api/pos/checkout'


@pytest.fixture
def smtp_box(monkeypatch, store_settings):
    store_settings('smtp', host='smtp.example.com', port=587, sender='shop@example.com')
    box = []

    class FakeSMTP:
        def send_message(self, msg):
            box.append(msg)

        def quit(self):
            pass
    monkeypatch.setattr(mail_service, '_connect', lambda cfg: FakeSMTP())
    return box


# ---------------------------------------------------------------------------
# 報表
# ---------------------------------------------------------------------------
def test_report_margin_and_breakdowns(app, admin_client, staff_client):
    iid = make_item(app, '報表毛利餐', price=200, stock=50, unit_cost=80, category='報表分類')
    checkout(staff_client, [line(iid, 2)], url=POS, cash_received=400)
    checkout(staff_client, [line(iid)], url=POS, payment_method='Credit Card', payment_info={'auth_code': 'RPT001'})
    cancelled = checkout(staff_client, [line(iid)], url=POS, cash_received=200).get_json()['order_id']
    admin_client.post(f'/api/orders/{cancelled}/status', json={'status': 'Cancelled'})
    with app.app_context():
        r = report_service.build_report(default_tenant_id(app), date.today(), date.today())
    item = next(i for i in r['items'] if i['name'] == '報表毛利餐')
    assert (item['qty'], item['revenue'], round(item['cost']), round(item['margin'])) == (3, 600, 240, 360)
    assert r['categories']['報表分類']['margin'] == 360
    assert r['cancelled_count'] >= 1 and r['net'] >= 600
    assert r['by_method']['信用卡']['in'] >= 200 and r['by_method']['現金']['out'] >= 200
    hour = datetime.now()
    assert r['heat'][hour.weekday()][hour.hour] >= 2
    assert r['by_source']['POS']['count'] >= 2 and 'pos_user' in r['by_cashier']

    page = admin_client.get('/admin/reports?preset=today').get_data(as_text=True)
    assert '報表毛利餐' in page and '毛利' in page
    xlsx = admin_client.get('/admin/reports/export?preset=today')
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.data))
    assert {'總覽', '每日', '品項', '分類', '付款方式', '點餐管道', '時段熱度', '收銀員'} <= set(wb.sheetnames)
    assert any(row[0].value == '報表毛利餐' for row in wb['品項'].iter_rows(min_row=2))
    assert staff_client.get('/admin/reports').status_code == 403
    assert admin_client.get('/admin/reports?start=2020-01-01&end=2026-01-01').status_code in (302, 400)


def test_scheduled_report_email(app, admin_client, store_settings, smtp_box):
    store_settings('report_email', enabled=True, frequency='daily', send_time='08:00',
                   recipients='boss@example.com\nnot-an-email')
    admin_client.post('/admin/reports/send')
    assert_flash(admin_client, '1 位收件者')
    msg = smtp_box[-1]
    assert msg['To'] == 'boss@example.com' and '營運報表' in msg['Subject']
    attachment = next(msg.iter_attachments())
    assert attachment.get_filename().endswith('.xlsx')
    # 每週：只在週一寄
    store_settings('report_email', frequency='weekly')
    tid = default_tenant_id(app)
    with app.app_context():
        monday = date.today() - timedelta(days=date.today().weekday())
        assert report_service.send_scheduled_report(tid, monday + timedelta(days=1)) == 0
        assert report_service.send_scheduled_report(tid, monday) == 1
        assert report_service.period_for('weekly', monday) == (monday - timedelta(days=7), monday - timedelta(days=1))


def test_daily_jobs_run_once_per_day(app, store_settings, smtp_box, monkeypatch):
    store_settings('report_email', enabled=True, frequency='daily', send_time='08:00', recipients='boss@example.com')
    store_settings('backup', enabled=True, time='03:30', keep_days=14)
    morning = datetime.now().replace(hour=7, minute=0)
    later = datetime.now().replace(hour=9, minute=0)
    with app.app_context():
        from app.sync.runtime import set_state
        for key in ('job:backup', f'job:report:{default_tenant_id(app)}', f'job:birthday:{default_tenant_id(app)}'):
            set_state(key, '')
        db.session.commit()
        before = len(smtp_box)
        done = maintenance_service.run_daily_jobs(morning)
        assert 'report' not in done and len(smtp_box) == before, '還沒到寄送時間'
        if db.engine.dialect.name == 'sqlite':
            assert 'backup' in done
        assert 'report' in maintenance_service.run_daily_jobs(later)
        assert maintenance_service.run_daily_jobs(later) == [], '同一天不重複執行'
        assert len(smtp_box) == before + 1


# ---------------------------------------------------------------------------
# 系統監控 / 備份
# ---------------------------------------------------------------------------
def test_system_page_and_backup(app, admin_client, manager_client):
    page = admin_client.get('/admin/system').get_data(as_text=True)
    assert '系統監控' in page or '健康' in page
    if 'sqlite' not in app.config['SQLALCHEMY_DATABASE_URI']:
        return
    assert manager_client.post('/admin/system/backup').status_code == 403
    admin_client.post('/admin/system/backup')
    msgs = assert_flash(admin_client, '已建立備份')
    name = msgs[-1].split('已建立備份 ')[-1].strip()
    assert name in admin_client.get('/admin/system').get_data(as_text=True)
    r = admin_client.get(f'/admin/system/backup/{name}')
    assert r.status_code == 200
    with zipfile.ZipFile(io.BytesIO(r.data)) as zf:
        assert zf.namelist() == ['menu.db'] and zf.read('menu.db')[:16] == b'SQLite format 3\x00'
    assert admin_client.get('/admin/system/backup/..%2Fmenu.db').status_code == 404
    assert admin_client.get('/admin/system/backup/other.zip').status_code == 404
    admin_client.post('/admin/system/test_alert')
    assert_flash(admin_client, '尚未設定任何通知管道')


def test_health_counts(app):
    with app.app_context():
        h = maintenance_service.health(default_tenant_id(app))
    assert h['db_ok'] is True and h['revision'] and {'invoice_pending', 'print_failed', 'refund_manual'} <= set(h)


# ---------------------------------------------------------------------------
# 打卡與班表
# ---------------------------------------------------------------------------
def test_timeclock_punch_and_admin(app, admin_client):
    login_staff(app, 'clock_user', 'staff')
    c = app.test_client()
    c.get('/?shop=default')
    assert c.get('/timeclock').status_code == 200
    c.post('/timeclock', data={'username': 'clock_user', 'password': 'wrong', 'action': 'in'})
    with app.app_context():
        assert TimeClock.query.count() == 0 or not TimeClock.query.filter(TimeClock.staff.has(username='clock_user')).count()
    c.post('/timeclock', data={'username': 'clock_user', 'password': 'StaffPass!2026', 'action': 'in'})
    assert_flash(c, '上班打卡')
    c.post('/timeclock', data={'username': 'clock_user', 'password': 'StaffPass!2026', 'action': 'in'})
    assert_flash(c, '請先下班打卡')
    c.post('/timeclock', data={'username': 'clock_user', 'password': 'StaffPass!2026', 'action': 'out'})
    assert_flash(c, '下班打卡')
    with app.app_context():
        rec = TimeClock.query.filter(TimeClock.staff.has(username='clock_user')).one()
        rid, staff_id = rec.id, rec.staff_account_id
    # 後台修正時間
    day = date.today().isoformat()
    admin_client.post('/admin/timeclock', data={'action': 'edit_record', 'record_id': rid,
                                                'clock_in': f'{day}T09:00', 'clock_out': f'{day}T17:30', 'note': '補登'})
    with app.app_context():
        rec = db.session.get(TimeClock, rid)
        assert rec.hours == 8.5 and rec.edited_by == 'admin_test'
    admin_client.post('/admin/timeclock', data={'action': 'edit_record', 'record_id': rid,
                                                'clock_in': f'{day}T18:00', 'clock_out': f'{day}T17:30'})
    assert_flash(admin_client, '不可早於')
    admin_client.post('/admin/timeclock', data={'action': 'add_schedule', 'staff_id': staff_id, 'work_date': day,
                                                'start_time': '10:00', 'end_time': '18:00', 'repeat_days': '2'})
    page = admin_client.get(f'/admin/timeclock?start={day}').get_data(as_text=True)
    assert 'clock_user' in page and '10:00' in page and '8.5' in page
    x = admin_client.get(f'/admin/timeclock/export?start={day}&end={day}')
    wb = load_workbook(io.BytesIO(x.data))
    assert any(r[0].value == 'clock_user' and r[3].value == 8.5 for r in wb.active.iter_rows(min_row=2))
    admin_client.post('/admin/timeclock', data={'action': 'add_schedule', 'staff_id': staff_id, 'work_date': day,
                                                'start_time': '25:00', 'end_time': '18:00'})
    assert_flash(admin_client, 'HH:MM')


# ---------------------------------------------------------------------------
# 出餐追蹤 / PWA
# ---------------------------------------------------------------------------
def test_track_page_and_pwa(app, staff_client):
    iid = make_item(app, '追蹤測試餐', price=45, stock=10)
    d = checkout(staff_client, [line(iid)], url=POS, cash_received=45).get_json()
    assert d['track_qr'].startswith('data:image/')
    c = app.test_client()
    page = c.get(f"/track/{d['order_uuid']}")
    assert page.status_code == 200 and '追蹤測試餐' in page.get_data(as_text=True)
    api = c.get(f"/api/track/{d['order_uuid']}").get_json()
    assert api['status'] == 'Pending' and api['pickup_number'] == d['pickup_number']
    staff_client.post(f"/api/orders/{d['order_id']}/status", json={'status': 'Completed'})
    assert c.get(f"/api/track/{d['order_uuid']}").get_json()['status'] == 'Completed'
    assert c.get('/track/short').status_code == 404
    assert c.get('/track/' + 'f' * 36).status_code == 404
    m = c.get('/manifest.webmanifest')
    assert m.mimetype == 'application/manifest+json' and m.get_json()['display'] == 'standalone'
    sw = c.get('/sw.js')
    assert sw.status_code == 200 and sw.headers['Service-Worker-Allowed'] == '/' and 'offline' in sw.get_data(as_text=True)
    assert c.get('/offline').status_code == 200
    home = guest_client(app).get('/').get_data(as_text=True)
    assert 'manifest.webmanifest' in home and 'pwa.js' in home


# ---------------------------------------------------------------------------
# 菜單進階設定 / 多語系 / 過敏原
# ---------------------------------------------------------------------------
def test_menu_pro_and_i18n(app, admin_client, store_settings):
    iid = make_item(app, '多語系牛肉麵', price=160, stock=20, category='多語系分類')
    other = make_item(app, '代碼重複測試', price=10, stock=5, external_code='DUP-1')
    form = {'action': 'item', 'item_id': iid, 'avail_days': list('1234567'), 'takeout_price': '150',
            'allergens': ['peanut', 'gluten', 'bogus'], 'diet': 'vegan', 'spicy_level': '2', 'name_en': 'Beef Noodle Soup',
            'name_ja': '牛肉麺', 'description_en': 'Slow braised', 'external_code': 'BN-01', 'unit_cost': '60'}
    admin_client.post('/admin/menu_pro', data=form)
    with app.app_context():
        item = db.session.get(MenuItem, iid)
        assert item.takeout_price == 150 and item.allergens == 'peanut,gluten' and item.diet == 'vegan'
        assert item.price_for('takeout') == 150 and item.price_for('dine_in') == 160
    admin_client.post('/admin/menu_pro', data=dict(form, external_code='DUP-1'))
    assert_flash(admin_client, '已被「代碼重複測試」使用')
    admin_client.post('/admin/menu_pro', data=dict(form, avail_days=[]))
    assert_flash(admin_client, '至少勾選一個供應星期')
    admin_client.post('/admin/menu_pro', data=dict(form, avail_start='9:00'))
    assert_flash(admin_client, 'HH:MM')
    admin_client.post('/admin/menu_pro', data={'action': 'categories', 'category': ['多語系分類'],
                                               'en__多語系分類': 'Noodles', 'ja__多語系分類': '麺類'})
    store_settings('menu', languages=['en', 'ja'], show_allergens=True)
    c = guest_client(app)
    en = c.get('/?lang=en').get_data(as_text=True)
    assert 'Beef Noodle Soup' in en and 'Noodles' in en and 'Add to cart' in en
    assert 'Clear' in en and '<built-in method clear' not in en
    assert 'Order with ease' in en and 'Your cart is empty' in en
    landing = app.test_client().get('/?lang=en').get_data(as_text=True)
    assert 'landingLangSwitch' in landing and 'Member login' in landing
    assert 'Beef Noodle Soup' in c.get('/').get_data(as_text=True), '語言記在 session'
    ja = c.get('/?lang=ja').get_data(as_text=True)
    assert '牛肉麺' in ja and 'カートに追加' in ja
    zh = c.get('/?lang=zh').get_data(as_text=True)
    assert '多語系牛肉麵' in zh
    store_settings('menu', languages=[])
    assert '多語系牛肉麵' in c.get('/?lang=en').get_data(as_text=True), '關閉的語言回到中文'
    assert admin_client.get('/admin/menu_pro').status_code == 200
