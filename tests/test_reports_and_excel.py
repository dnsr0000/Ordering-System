"""報表與營運分析、Excel 匯出 / 匯入 (含公式注入防護)、儀表板範圍與稽核紀錄"""
import io
import os
import re
import zipfile
from datetime import datetime, timedelta

import openpyxl
import pytest

from app.extensions import db
from app.models.menu import MenuItem
from app.models.order import Order, OrderItem
from app.models.user import StaffAuditLog, User
from app.services import analytics_service
from app.services.analytics_service import (
    DEFAULT_MINUTES_PER_ORDER,
    build_order_analytics,
    mean_without_outliers,
)
from tests.support import (
    assert_flash,
    default_tenant_id,
    make_member,
)

EXPECTED_SHEETS = {'會員資料(User)', '菜單品項(MenuItem)', '優惠券(Coupon)', '會員持券(UserCoupon)', '訂單總覽(Order)',
                   '訂單明細(OrderItem)', '區間營運總覽(Analytics)', '區間熱銷排行(ItemSales)',
                   '管理員進出日誌(AdminLog)', '員工操作稽核(StaffAuditLog)'}


def import_file(client, content, filename):
    data = content.encode('utf-8') if isinstance(content, str) else content
    return client.post('/admin/import_smart', data={'file': (io.BytesIO(data), filename)}, content_type='multipart/form-data')


def get_item(app, name):
    with app.app_context():
        item = MenuItem.query.filter_by(name=name).first()
        db.session.expunge_all()
        return item


@pytest.fixture
def cleanup_items(app):
    names = []
    yield names
    with app.app_context():
        MenuItem.query.filter(MenuItem.name.in_(names)).delete(synchronize_session=False)
        db.session.commit()


def test_analytics_and_ai_fallback(app, admin_client):
    """沒有 Gemini 金鑰時，營運建議與菜單文案改用本地規則引擎產生"""
    d = admin_client.get('/api/admin_live_orders').get_json()
    assert len(d['analytics']['ai_advice']) > 10
    r = admin_client.post('/api/generate_item_description', json={'name': '招牌和牛堡', 'category': '主餐'})
    assert r.status_code == 200 and r.get_json()['description']


def test_excel_export(app, admin_client, ctx):
    r = admin_client.post('/admin/export_excel', data={'tables': ['User', 'MenuItem', 'Coupon', 'UserCoupon', 'Order',
                                                                 'OrderItem', 'Analytics', 'AdminLog']})
    assert r.status_code == 200
    ctx['exported_xlsx'] = r.data
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert EXPECTED_SHEETS <= set(wb.sheetnames)
    header = [c.value for c in wb['訂單總覽(Order)'][2]]
    for col in ['訂單來源(Source)', '收銀員(Cashier)', '應付總額(Total)', '實收金額(Received)', '找零金額(Change)']:
        assert col in header
    assert wb['員工操作稽核(StaffAuditLog)'].max_row >= 5


def test_import_csv_and_excel_upsert(app, admin_client, ctx):
    r = import_file(admin_client, "餐點名稱,分類,單價,客製化群組,商品描述\n宇治抹茶拿鐵,飲料,65,ice_sugar,京都宇治抹茶鮮奶\n", 'new_menu_items.csv')
    assert r.status_code == 302
    wb = openpyxl.load_workbook(io.BytesIO(ctx['exported_xlsx']))
    ws = wb['菜單品項(MenuItem)']
    headers = [c.value for c in ws[1]]
    name_col, price_col = headers.index('餐點名稱(Name)') + 1, headers.index('單價(Price)') + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=name_col).value == '錫蘭紅茶':
            ws.cell(row=row, column=price_col).value = 45
    buf = io.BytesIO()
    wb.save(buf)
    assert import_file(admin_client, buf.getvalue(), 'updated_store_data.xlsx').status_code == 302
    assert_flash(admin_client, '智慧匯入完成')
    assert get_item(app, '宇治抹茶拿鐵').price == 65
    assert get_item(app, '錫蘭紅茶').price == 45


def test_import_keeps_stock_without_stock_column(app, admin_client, cleanup_items):
    """檔案沒有庫存欄位時，既有品項的庫存與售完狀態不被覆蓋"""
    cleanup_items += ['庫存保留測試茶', '庫存保留測試茶二號']
    import_file(admin_client, "餐點名稱,分類,單價,庫存\n庫存保留測試茶,飲料,40,7\n", 'a.csv')
    assert get_item(app, '庫存保留測試茶').stock == 7
    with app.app_context():
        MenuItem.query.filter_by(name='庫存保留測試茶').update({'stock': 3})
        db.session.commit()
    import_file(admin_client, "餐點名稱,分類,單價\n庫存保留測試茶,飲料,55\n庫存保留測試茶二號,飲料,35\n", 'b.csv')
    tea = get_item(app, '庫存保留測試茶')
    assert (tea.price, tea.stock) == (55, 3) and get_item(app, '庫存保留測試茶二號').stock == 50
    with app.app_context():
        MenuItem.query.filter_by(name='庫存保留測試茶').update({'stock': 0, 'is_sold_out': True})
        db.session.commit()
    import_file(admin_client, "餐點名稱,分類,單價\n庫存保留測試茶,飲料,60\n", 'c.csv')
    tea = get_item(app, '庫存保留測試茶')
    assert (tea.price, tea.stock, tea.is_sold_out) == (60, 0, True)
    import_file(admin_client, "餐點名稱,分類,單價,庫存\n庫存保留測試茶,飲料,60,12\n", 'd.csv')
    tea = get_item(app, '庫存保留測試茶')
    assert (tea.stock, tea.is_sold_out) == (12, False)


def test_import_is_atomic(app, admin_client, cleanup_items):
    """任一工作表解析失敗時整批不寫入"""
    cleanup_items += ['原子匯入測試餐A', '原子匯入測試餐B']

    def workbook(with_bad_sheet):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '菜單品項(MenuItem)'
        ws.append(['餐點名稱(Name)', '分類(Category)', '單價(Price)'])
        ws.append(['原子匯入測試餐A', '點心', 30])
        ws2 = wb.create_sheet('會員資料(User)')
        ws2.append(['姓名(Name)', '電話(Phone)', '紅利點數(Points)'])
        ws2.append(['原子匯入測試會員', '0900111222', 10])
        if with_bad_sheet:
            ws3 = wb.create_sheet('菜單_壞資料')
            ws3.append(['品名', '金額'])
            ws3.append(['原子匯入測試餐B', 40])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    import_file(admin_client, workbook(True), 'atomic.xlsx')
    assert_flash(admin_client, '解析失敗')
    with app.app_context():
        assert MenuItem.query.filter(MenuItem.name.like('原子匯入測試餐%')).count() == 0
        assert User.query.filter_by(phone='0900111222').count() == 0
    import_file(admin_client, workbook(False), 'atomic.xlsx')
    assert get_item(app, '原子匯入測試餐A').price == 30
    with app.app_context():
        member = User.query.filter_by(phone='0900111222').first()
        assert member.points == 10
        db.session.delete(member)
        db.session.commit()


def test_import_reward_guards(app, admin_client, cleanup_items):
    """匯入紅利防呆：點數為 0 或高於單價關閉紅利、特惠點數不合法歸 0"""
    names = ['紅利防呆測試A', '紅利防呆測試B', '紅利防呆測試C', '紅利防呆測試D']
    cleanup_items += names
    header = "餐點名稱,分類,單價,是否開放紅利兌換,兌換點數,特惠點數\n"
    import_file(admin_client, header + "紅利防呆測試A,點心,50,是,80,0\n紅利防呆測試B,點心,50,是,30,40\n"
                                       "紅利防呆測試C,點心,50,是,30,10\n紅利防呆測試D,點心,50,是,0,0\n", 'r.csv')
    got = {n: (i.is_reward, i.reward_points, i.reward_discount_points) for n in names for i in [get_item(app, n)]}
    assert got == {'紅利防呆測試A': (False, 0, 0), '紅利防呆測試B': (True, 30, 0),
                   '紅利防呆測試C': (True, 30, 10), '紅利防呆測試D': (False, 0, 0)}
    import_file(admin_client, header + "紅利防呆測試C,點心,50,是,99,0\n", 'r2.csv')
    c = get_item(app, '紅利防呆測試C')
    assert (c.is_reward, c.reward_points, c.reward_discount_points) == (False, 0, 0)


def test_import_total_stock_never_below_stock(app, admin_client, cleanup_items):
    cleanup_items.append('總量修正測試茶')

    def run(csv):
        assert import_file(admin_client, csv, 't.csv').status_code == 302
        it = get_item(app, '總量修正測試茶')
        return it.stock, it.total_stock

    assert run("餐點名稱,分類,單價,庫存,總量\n總量修正測試茶,飲料,30,10,10\n") == (10, 10)
    assert run("餐點名稱,分類,單價,庫存\n總量修正測試茶,飲料,30,25\n") == (25, 25)
    assert run("餐點名稱,分類,單價,總量\n總量修正測試茶,飲料,30,5\n") == (25, 25)
    assert run("餐點名稱,分類,單價,庫存,總量\n總量修正測試茶,飲料,30,8,20\n") == (8, 20)
    assert run("餐點名稱,分類,單價,總量\n總量修正測試茶,飲料,30,40\n") == (8, 40)


BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def test_excel_formula_injection_and_import_points(app, admin_client):
    make_member(app, '=1+1+M15', '0933150101', points=3)
    r = admin_client.post('/admin/export_excel', data={'tables': 'User'})
    ws = openpyxl.load_workbook(io.BytesIO(r.data))['會員資料(User)']
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=1+1+M15')]
    assert cells and all(c.data_type == 's' for c in cells)

    def workbook(rows, with_points=True):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '會員資料(User)'
        ws.append(['姓名(Name)', '電話(Phone)'] + (['紅利點數(Points)'] if with_points else []))
        for row in rows:
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    for buf in (workbook([['M15零點', '0933150011', 0], ['M15有點', '0933150012', 7]]),
                workbook([['M15無點數欄', '0933150013']], with_points=False)):
        r = admin_client.post('/admin/import_smart', data={'file': (buf, 'u.xlsx')}, content_type='multipart/form-data')
        assert r.status_code == 302
    with app.app_context():
        pts = {u.phone: u.points for u in User.query.filter(User.phone.in_(['0933150011', '0933150012', '0933150013']))}
    assert pts == {'0933150011': 0, '0933150012': 7, '0933150013': 20}


def test_dashboard_scope_and_audits(app, admin_client, kitchen_client):
    tid = default_tenant_id(app)
    old = datetime.now() - timedelta(days=3)
    with app.app_context():
        db.session.add_all([Order(tenant_id=tid, table_number='M15舊單甲', total_price=11, status='Completed', pickup_number=1, created_at=old),
                            Order(tenant_id=tid, table_number='M15舊單乙', total_price=22, status='Pending', pickup_number=2, created_at=old)])
        db.session.commit()
    data = admin_client.get('/api/admin_live_orders').get_json()
    names = [o['user_name'] for o in data['orders']]
    assert 'M15舊單甲' not in names and 'M15舊單乙' not in names
    assert data['analytics']['stale_pending_count'] >= 1
    admin_client.post('/admin/add_coupon', data={'code': 'M15AUD', 'title': '稽核測試', 'discount_type': 'fixed', 'discount_value': '5'})
    kitchen_client.post('/api/kitchen_complete_all')
    with app.app_context():
        kinds = {a.action_type for a in StaffAuditLog.query.filter_by(tenant_id=tid)}
    assert {'COUPON_CHANGE', 'KITCHEN_COMPLETE_ALL'} <= kinds


def _order(status, minutes_to_done=None, age=30, **kw):
    now = datetime.now()
    created = now - timedelta(minutes=age)
    done = created + timedelta(minutes=minutes_to_done) if minutes_to_done is not None else None
    return Order(total_price=kw.pop('total_price', 10), status=status, pickup_number=1, created_at=created, completed_at=done, **kw)


def test_analytics_fields(app, admin_client):
    olds = [_order('Pending', age=60 * 48) for _ in range(5)]
    today = [_order('Completed', 6, age=10, total_price=100), _order('PickedUp', 2, age=8, total_price=0),
             _order('Pending', age=1, total_price=50)]
    a = build_order_analytics(olds + today)
    assert (a['pending_count'], a['eta_minutes'], a['stale_pending_count']) == (1, 4, 5)
    assert (a['completed_today'], a['avg_order_time'], a['avg_order_value']) == (2, 4.0, 75.0)
    live = admin_client.get('/api/admin_live_orders').get_json()['analytics']
    assert {'completed_today', 'avg_order_time', 'stale_pending_count'} <= set(live)


def test_eta_excludes_outliers():
    m, n = mean_without_outliers([4, 4, 5, 5, 4, 60])
    assert n == 5 and round(m, 1) == 4.4
    assert mean_without_outliers([4, 5, 200]) == (4.5, 2)
    assert mean_without_outliers([]) == (0.0, 0)
    orders = [_order('Completed', d) for d in (4, 4, 5, 5, 4)] + [_order('Completed', 120), _order('Completed', None)]
    orders += [_order('Pending', age=1) for _ in range(3)]
    a = build_order_analytics(orders)
    assert (a['avg_order_time'], a['eta_based_on_samples'], a['pending_count'], a['eta_minutes'], a['eta_per_order_minutes']) == (4.4, 5, 3, 14, 4.4)
    a2 = build_order_analytics([_order('Pending', age=1), _order('Pending', age=2)])
    assert a2['eta_based_on_samples'] == 0 and a2['eta_minutes'] == int(2 * DEFAULT_MINUTES_PER_ORDER)
    assert build_order_analytics([_order('Completed', 4)])['eta_minutes'] == 0


def test_top_items_with_revenue(app, admin_client):
    def with_items(items, status='Completed'):
        o = _order(status, 4, age=5)
        o.items = [OrderItem(item_name=n, price=p, quantity=q, is_reward=r) for n, p, q, r in items]
        return o
    orders = [with_items([('香檸氣泡飲', 80, 1, False), ('紅茶', 0, 1, True)]),
              with_items([('香檸氣泡飲', 80, 2, False), ('雞排', 70, 1, False)]),
              with_items([('雞排', 70, 1, False)], status='Cancelled')]
    a = build_order_analytics(orders, limit=5)
    top = {t['name']: t for t in a['top_items']}
    assert (top['香檸氣泡飲']['quantity'], top['香檸氣泡飲']['revenue']) == (3, 240)
    assert (top['雞排']['quantity'], top['雞排']['revenue'], top['紅茶']['revenue']) == (1, 70, 0)
    assert a['top_items'][0]['name'] == '香檸氣泡飲' and len(build_order_analytics(orders, limit=1)['top_items']) == 1

    # 前端讀取的 analytics.<欄位> 必須都由後端提供 (模板與靜態 JS 都檢查)
    sources = [open(os.path.join(BASE_DIR, 'templates', 'admin.html'), encoding='utf-8').read()]
    js_dir = os.path.join(BASE_DIR, 'static', 'js')
    if os.path.isdir(js_dir):
        sources += [open(os.path.join(js_dir, f), encoding='utf-8').read() for f in os.listdir(js_dir) if f.endswith('.js')]
    used = set().union(*(re.findall(r'analytics\.(\w+)', s) for s in sources))
    assert used <= set(a) | {'ai_advice'}, sorted(used - set(a))
    live = admin_client.get('/api/admin_live_orders?limit=2').get_json()['analytics']
    assert len(live['top_items']) <= 2 and all('revenue' in t for t in live['top_items'])
    assert 'id="stat-top-items"' in admin_client.get('/admin').get_data(as_text=True)
    assert analytics_service.load_dashboard_orders  # 儀表板與即時 API 共用同一個查詢


def test_import_rejects_xlsx_bomb(app, admin_client):
    """xlsx 是 zip：解壓後超過上限的檔案在解析前就拒絕 (避免小檔案耗盡記憶體)"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', '<Types/>')
        zf.writestr('xl/worksheets/sheet1.xml', ' ' * (31 * 1024 * 1024))
    buf.seek(0)
    r = admin_client.post('/admin/import_smart', data={'file': (buf, 'bomb.xlsx')}, content_type='multipart/form-data')
    assert r.status_code == 302
    assert_flash(admin_client, '檔案內容過大')
