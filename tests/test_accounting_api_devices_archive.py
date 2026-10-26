"""第三輪 #13~#16：會計匯出、對外 API / Webhook、裝置管理、舊資料封存"""
import hashlib
import hmac
import json
import re
from contextlib import contextmanager
from datetime import date, datetime, timedelta

import pytest
from sqlalchemy import update as sa_update

from app.extensions import db, limiter
from app.models.integration import ApiKey, Device, WebhookDelivery, WebhookEndpoint
from app.models.menu import MenuItem
from app.models.order import Order
from app.models.payment import Invoice, PaymentRecord
from app.models.sync import SyncTombstone
from app.models.tenant import Tenant
from app.models.user import CustomerLog, User
from app.services import monitor_service, webhook_service
from tests.support import assert_flash, checkout, line, login_staff, make_item


@pytest.fixture
def no_rate_limit():
    limiter.enabled = False
    yield
    limiter.enabled = True


@pytest.fixture
def alerts(monkeypatch):
    out = []
    monkeypatch.setattr(monitor_service, 'alert', lambda msg, **kw: out.append(msg))
    return out


def _tenant(app, slug):
    with app.app_context():
        t = Tenant.query.filter_by(slug=slug).first()
        if t is None:
            t = Tenant(name=f'測試 {slug}', slug=slug)
            db.session.add(t)
            db.session.commit()
        return t.id


@contextmanager
def _ctx(app, tenant_id):
    from app.tenant_context import _use_tenant
    with app.test_request_context('/'):
        _use_tenant(db.session.get(Tenant, tenant_id), remember=False)
        yield


def _item(tid, name, price=100):
    item = MenuItem(tenant_id=tid, name=name, price=price, stock=999, total_stock=999)
    db.session.add(item)
    db.session.commit()
    return item.id


def _buy(tid, item_id, qty=1, **kw):
    from app.services.order_service import process_checkout
    order, _ = process_checkout(tenant_id=tid, items=[{'menu_item_id': item_id, 'quantity': qty}], **kw)
    db.session.commit()
    return order


# ===========================================================================
# 13. 會計匯出
# ===========================================================================
def test_accounting_journal_balances_and_exports(app):
    from app.services import accounting_service
    tid = _tenant(app, 'acct-shop')
    day = date(2026, 9, 15)
    at = datetime(2026, 9, 15, 12, 0)
    with _ctx(app, tid):
        for kind, method, amount in (('sale', 'Cash', 1050), ('settle', 'Credit Card', 2100), ('sale', '儲值金', 210),
                                     ('refund', 'Cash', -105), ('topup', 'Cash', 1000), ('gift_sale', 'Credit Card', 500),
                                     ('pay_out', 'Cash', -100), ('pay_in', 'Cash', 300)):
            db.session.add(PaymentRecord(tenant_id=tid, kind=kind, method=method, amount=amount, created_at=at))
        db.session.add_all([
            Invoice(tenant_id=tid, relate_number='ACCT-1', status='issued', amount=1050, invoice_no='AB12345678',
                    invoice_date=at, buyer_identifier='12345675', carrier_type='company'),
            Invoice(tenant_id=tid, relate_number='ACCT-2', status='void', amount=210, invoice_no='AB12345679',
                    invoice_date=at, carrier_type='print'),
        ])
        db.session.commit()
        vouchers = accounting_service.journal(tid, day, day)
        assert len(vouchers) == 1
        v = vouchers[0]
        assert v['no'] == 'JV20260915' and v['debit'] == v['credit'] == 5365
        lines = {(ln['code'], 'D' if ln['debit'] else 'C'): ln['debit'] or ln['credit'] for ln in v['lines']}
        assert lines[('1111', 'D')] == 2350 and lines[('1111', 'C')] == 205        # 現金
        assert lines[('1141', 'D')] == 2600                                       # 信用卡
        assert lines[('2151', 'D')] == 210 and lines[('2151', 'C')] == 1000       # 儲值金：沖減 / 預收
        assert lines[('2152', 'C')] == 500                                        # 禮物卡預收
        assert lines[('4111', 'C')] == 3200 and lines[('2204', 'C')] == 160       # 3360 含稅
        assert lines[('4171', 'D')] == 100 and lines[('2204', 'D')] == 5          # 退回 105 含稅
        assert lines[('6188', 'D')] == 100 and lines[('1112', 'C')] == 300
        inv = accounting_service.invoices(tid, day, day)
        assert inv[0][:8] == ['AB12345678', '2026-09-15', '12345675', '', 1000, 50, 1050, '開立']
        assert inv[1][4:8] == [0, 0, 0, '作廢']
    owner = login_staff(app, 'acct_owner', 'owner', tenant_id=tid)
    page = owner.get('/admin/accounting?start=2026-09-15&end=2026-09-15').get_data(as_text=True)
    assert 'data-voucher="JV20260915"' in page and '借 $5,365 / 貸 $5,365' in page and '發票明細 CSV (2 張)' in page
    csv_resp = owner.get('/admin/accounting?start=2026-09-15&end=2026-09-15&export=journal')
    text = csv_resp.data.decode('utf-8')
    assert text.startswith('﻿傳票日期,傳票號碼') and '2026-09-15,JV20260915,1,' in text
    assert owner.get('/admin/accounting?start=2026-09-15&end=2026-09-15&export=invoices&format=xlsx').data[:2] == b'PK'
    assert login_staff(app, 'acct_staff', 'staff', tenant_id=tid).get('/admin/accounting').status_code == 403


# ===========================================================================
# 14. 對外 API / Webhook
# ===========================================================================
def _create_key(owner, scopes, name='ERP'):
    page = owner.post('/admin/integrations', data={'action': 'create_key', 'name': name, 'scopes': scopes},
                      follow_redirects=True).get_data(as_text=True)
    return re.search(r'<code class="user-select-all fs-6">(kk_[A-Za-z0-9_-]+)</code>', page).group(1)


def test_public_api_scopes_and_tenant_isolation(app, no_rate_limit):
    tid = _tenant(app, 'api-shop')
    owner = login_staff(app, 'api_owner', 'owner', tenant_id=tid)
    key = _create_key(owner, ['orders:read', 'menu:read', 'members:read', 'reports:read'])
    writer = _create_key(owner, ['menu:write'], name='庫存同步')
    with _ctx(app, tid):
        item_id = _item(tid, 'API 測試餐', 120)
        db.session.add(User(tenant_id=tid, name='API 會員', phone='0977400001', photo_path='', points=15))
        db.session.commit()
        first = _buy(tid, item_id, 2, now=datetime.now() - timedelta(minutes=10))
        second = _buy(tid, item_id, 1)
        uuids = (first.uuid, second.uuid)
        second_day = second.created_at.date()
        same_day = 2 if first.created_at.date() == second_day else 1
    other = make_item(app, 'API 別店餐', price=50)
    other_order = checkout(app.test_client(), [line(other)])   # 預設店家的訂單 (不可被讀到)
    other_uuid = None
    if other_order.status_code == 200:
        with app.app_context():
            other_uuid = db.session.get(Order, other_order.get_json()['order_id']).uuid
    c = app.test_client()
    h = {'Authorization': f'Bearer {key}'}
    assert c.get('/api/v1/orders').status_code == 401
    assert c.get('/api/v1/orders', headers={'Authorization': 'Bearer kk_wrong'}).status_code == 401
    data = c.get('/api/v1/orders', headers=h).get_json()
    assert [o['uuid'] for o in data['orders']] == list(uuids) and data['orders'][0]['items'][0]['quantity'] == 2
    page1 = c.get('/api/v1/orders?limit=1', headers=h).get_json()
    assert page1['orders'][0]['uuid'] == uuids[0] and page1['next_since']
    page2 = c.get('/api/v1/orders', query_string=dict(page1['next'], limit=1), headers=h).get_json()
    assert [o['uuid'] for o in page2['orders']] == [uuids[1]]
    # 異動追蹤 (cursor)：新增與狀態改變依提交順序
    feed = c.get('/api/v1/orders?cursor=0', headers=h).get_json()
    assert [o['uuid'] for o in feed['orders']] == list(uuids)
    assert c.get(f"/api/v1/orders?cursor={feed['next_cursor']}", headers=h).get_json()['orders'] == []
    with _ctx(app, tid):
        o = Order.query.filter_by(uuid=uuids[0]).one()
        o.status = 'Completed'
        db.session.commit()
    changed = c.get(f"/api/v1/orders?cursor={feed['next_cursor']}", headers=h).get_json()
    assert [(o['uuid'], o['status']) for o in changed['orders']] == [(uuids[0], 'Completed')]
    assert c.get(f'/api/v1/orders/{uuids[1]}', headers=h).get_json()['order']['total_price'] == 120
    if other_uuid:
        assert c.get(f'/api/v1/orders/{other_uuid}', headers=h).status_code == 404
    menu = c.get('/api/v1/menu', headers=h).get_json()['items']
    assert [m['name'] for m in menu] == ['API 測試餐']
    assert c.post(f'/api/v1/menu/{item_id}/availability', json={'sold_out': True}, headers=h).status_code == 403
    w = {'Authorization': f'Bearer {writer}'}
    r = c.post(f'/api/v1/menu/{item_id}/availability', json={'sold_out': True, 'stock': 5}, headers=w)
    assert r.get_json()['item'] == {'id': item_id, 'stock': 5, 'is_sold_out': True}
    assert c.post(f'/api/v1/menu/{other}/availability', json={'sold_out': True}, headers=w).status_code == 404
    assert c.get('/api/v1/menu', headers=w).status_code == 403
    m = c.get('/api/v1/members/0977400001', headers=h).get_json()['member']
    assert (m['name'], m['points']) == ('API 會員', 15)
    day = second_day
    daily = c.get(f'/api/v1/reports/daily?date={day}', headers=h).get_json()
    assert daily['orders'] == same_day and daily['net'] == (360 if same_day == 2 else 120)   # 跨午夜時只算當天的
    page = owner.get('/admin/integrations').get_data(as_text=True)
    assert 'ERP' in page and key[:10] in page and key not in page          # 只顯示一次
    with app.app_context():
        kid = ApiKey.query.filter_by(name='ERP', tenant_id=tid).one().id
        assert db.session.get(ApiKey, kid).last_used_at is not None
    owner.post('/admin/integrations', data={'action': 'revoke_key', 'key_id': kid})
    assert c.get('/api/v1/orders', headers=h).status_code == 401


def test_webhook_events_signature_retry_and_ssrf_guard(app, monkeypatch, alerts):
    tid = _tenant(app, 'hook-shop')
    owner = login_staff(app, 'hook_owner', 'owner', tenant_id=tid)
    r = owner.post('/admin/integrations', data={'action': 'create_hook', 'name': 'ERP', 'url': 'http://erp.example.test/x',
                                                'events': ['order.created']}, follow_redirects=True)
    assert 'https://' in r.get_data(as_text=True)
    owner.post('/admin/integrations', data={'action': 'create_hook', 'name': 'ERP', 'url': 'https://erp.example.test/hook',
                                            'events': list(webhook_service.WEBHOOK_EVENTS)})
    monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host: {'93.184.216.34'})
    sent = []

    def ok_post(url, body, headers):
        sent.append((url, body, headers))
        return 200

    with _ctx(app, tid):
        ep = WebhookEndpoint.query.filter_by(tenant_id=tid).one()
        item_id = _item(tid, 'Webhook 測試餐', 80)
        order = _buy(tid, item_id)
        assert webhook_service.scan(tid) == 1
        assert webhook_service.scan(tid) == 0                                # 沒有變化不重送
        assert webhook_service.deliver_due(tid, post=ok_post) == 1
        url, body, headers = sent[-1]
        payload = json.loads(body)
        assert url == ep.url and payload['event'] == 'order.created' and payload['data']['uuid'] == order.uuid
        expect = 'sha256=' + hmac.new(ep.secret.encode(), f"{headers['X-Kiosk-Timestamp']}.{body}".encode(),
                                      hashlib.sha256).hexdigest()
        assert headers['X-Kiosk-Signature'] == expect and headers['X-Kiosk-Delivery'] == payload['id']
        # 狀態改變 / 取消
        order = db.session.get(Order, order.id)
        order.status = 'Completed'
        db.session.commit()
        webhook_service.scan(tid)
        order.status = 'Cancelled'
        db.session.commit()
        webhook_service.scan(tid)
        webhook_service.deliver_due(tid, post=ok_post)
        assert [json.loads(b)['event'] for _u, b, _h in sent[-2:]] == ['order.status_changed', 'order.cancelled']
        # 線上待付款：付款完成時才送 created + paid
        online = Order(tenant_id=tid, total_price=80, status='Pending', payment_status='pending', payment_method='Credit Card',
                       order_type='外帶', pickup_number=0)
        db.session.add(online)
        db.session.commit()
        assert webhook_service.scan(tid) == 0
        online.payment_status = ''
        db.session.commit()
        assert webhook_service.scan(tid) == 2
        webhook_service.deliver_due(tid, post=ok_post)
        assert [json.loads(b)['event'] for _u, b, _h in sent[-2:]] == ['order.created', 'order.paid']

        # 失敗重試 → 6 次後放棄並通知
        webhook_service.queue_ping(ep)
        db.session.commit()
        now = datetime.now()
        for i in range(webhook_service.MAX_ATTEMPTS):
            webhook_service.deliver_due(tid, now=now, post=lambda u, b, h: 500)
            now += timedelta(minutes=200)
        d = WebhookDelivery.query.filter_by(endpoint_id=ep.id, event='ping').one()
        assert (d.status, d.attempts, d.last_error) == ('failed', 6, 'HTTP 500')
        assert any('ping 送出 6 次都失敗' in a for a in alerts)
        # 內網位址一律拒絕
        monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host: {'10.0.0.5'})
        webhook_service.queue_ping(ep)
        db.session.commit()
        webhook_service.deliver_due(tid, post=ok_post)
        d = WebhookDelivery.query.filter_by(endpoint_id=ep.id, event='ping').order_by(WebhookDelivery.id.desc()).first()
        assert d.status == 'pending' and '不可送到內網' in d.last_error
    page = owner.get('/admin/integrations').get_data(as_text=True)
    assert 'id="deliveryTable"' in page and '已送達' in page and '失敗' in page


# ===========================================================================
# 15. 裝置管理
# ===========================================================================
def test_device_pair_heartbeat_lock_and_reload(app, admin_client, no_rate_limit):
    r = admin_client.post('/admin/devices', data={'action': 'create', 'name': '測試點餐機', 'kind': 'kiosk'},
                          follow_redirects=True)
    assert '配對碼' in r.get_data(as_text=True)
    with app.app_context():
        d = Device.query.filter_by(name='測試點餐機').one()
        code, did = d.pair_code, d.id
    kiosk = app.test_client()
    kiosk.get('/?shop=default')
    kiosk.post('/device/pair', data={'code': '000000' if code != '000000' else '111111'})
    assert_flash(kiosk, '配對碼錯誤')
    r = kiosk.post('/device/pair', data={'code': code})
    assert r.status_code == 302 and r.headers['Location'].endswith('/')
    assert 'kiosk_device=kd_' in r.headers['Set-Cookie'] and 'HttpOnly' in r.headers['Set-Cookie']
    kiosk.post('/guest_login')
    html = kiosk.get('/').get_data(as_text=True)
    assert 'data-device="1"' in html and 'js/device.js' in html
    hb = kiosk.post('/api/device/heartbeat', json={'version': re.search(r'data-app-version="(\w+)"', html).group(1),
                                                  'path': '/'}).get_json()
    assert hb['registered'] and hb['version_ok'] and not hb['locked']
    item = make_item(app, '裝置鎖定測試餐', price=40, stock=10)
    assert checkout(kiosk, [line(item)]).status_code == 200
    # 遠端鎖定：畫面鎖定 + 操作拒絕
    admin_client.post('/admin/devices', data={'action': 'lock', 'device_id': did, 'message': '設備維修中'})
    hb = kiosk.post('/api/device/heartbeat', json={'version': 'old', 'path': '/'}).get_json()
    assert hb['locked'] and hb['message'] == '設備維修中' and not hb['version_ok']
    blocked = checkout(kiosk, [line(item)])
    assert blocked.status_code == 423 and '設備維修中' in blocked.get_json()['message']
    assert checkout(app.test_client(), [line(item)]).status_code in (200, 400, 401)   # 其他裝置不受影響
    page = admin_client.get('/admin/devices').get_data(as_text=True)
    assert f'data-device-id="{did}" data-online="1"' in page and '已鎖定' in page and '需重新整理' in page
    admin_client.post('/admin/devices', data={'action': 'unlock', 'device_id': did})
    admin_client.post('/admin/devices', data={'action': 'reload', 'device_id': did})
    hb = kiosk.post('/api/device/heartbeat', json={'version': 'old'}).get_json()
    assert hb['reload'] is True and not hb['locked']
    assert kiosk.post('/api/device/heartbeat', json={}).get_json()['reload'] is False   # 只重新整理一次
    assert checkout(kiosk, [line(item)]).status_code == 200
    admin_client.post('/admin/devices', data={'action': 'delete', 'device_id': did})
    assert kiosk.post('/api/device/heartbeat', json={}).get_json()['registered'] is False
    assert app.test_client().post('/api/device/heartbeat', json={}).get_json()['registered'] is False


# ===========================================================================
# 16. 舊資料封存
# ===========================================================================
def test_archive_moves_old_orders_and_logs_keeps_member_spend(app):
    import sqlalchemy as sa
    from app.services import archive_service, wallet_service
    tid = _tenant(app, 'archive-shop')
    old = datetime.now() - timedelta(days=3 * 365)
    with _ctx(app, tid):
        item_id = _item(tid, '封存測試餐', 100)
        u = User(tenant_id=tid, name='封存會員', phone='0977500001', photo_path='')
        db.session.add(u)
        db.session.commit()
        uid = u.id
        old_orders = [_buy(tid, item_id, 2, user_id=uid, now=old + timedelta(days=i)) for i in range(3)]
        for o in old_orders:
            o.status = 'PickedUp'
        stuck = _buy(tid, item_id, 1, now=old)
        stuck.payment_status = 'refund_required'    # 尚待退款：不封存
        recent = _buy(tid, item_id, 1, user_id=uid)
        db.session.add(CustomerLog(tenant_id=tid, user_type='guest', identifier='x', login_at=old))
        db.session.commit()
        old_ids = [o.id for o in old_orders]
        stuck_id, recent_id = stuck.id, recent.id
        before_spent = wallet_service.recompute_spent(db.session.get(User, uid))
        assert before_spent == 700
        tomb_before = SyncTombstone.query.count()
        preview = archive_service.preview(tid, archive_service.cutoff_for(tid))
        assert preview['order'] == 3 and preview['order_item'] == 3 and preview['customer_log'] >= 1
    owner = login_staff(app, 'archive_owner', 'owner', tenant_id=tid)
    page = owner.get('/admin/archive').get_data(as_text=True)
    assert 'data-archive-count="order">3<' in page and 'id="archiveRunBtn"' in page
    r = owner.post('/admin/archive', data={'action': 'run'}, follow_redirects=True)
    assert '訂單 3 筆' in r.get_data(as_text=True)
    with _ctx(app, tid):
        assert Order.query.filter(Order.id.in_(old_ids)).count() == 0
        assert db.session.get(Order, stuck_id) is not None and db.session.get(Order, recent_id) is not None
        assert PaymentRecord.query.filter(PaymentRecord.order_id.in_(old_ids)).count() == 0
        assert SyncTombstone.query.count() == tomb_before                  # 不產生同步刪除
        user = db.session.get(User, uid)
        assert user.archived_spent == 600 and wallet_service.recompute_spent(user) == before_spent
        assert archive_service.preview(tid, archive_service.cutoff_for(tid))['order'] == 0
        path = archive_service.archive_path(tid)
    engine = sa.create_engine('sqlite:///' + path)
    with engine.connect() as conn:
        assert conn.execute(sa.text('SELECT COUNT(*) FROM "order" WHERE tenant_id = :t'), {'t': tid}).scalar() == 3
        assert conn.execute(sa.text('SELECT COUNT(*) FROM customer_log')).scalar() >= 1
    engine.dispose()
    found = owner.get(f'/admin/archive?q={old_ids[0]}').get_data(as_text=True)
    assert 'id="archiveResults"' in found and '封存測試餐 x2' in found
    assert '封存測試餐' in owner.get('/admin/archive?q=0977500001').get_data(as_text=True)
    assert owner.get('/admin/archive?download=1').data[:15] == b'SQLite format 3'
    r = owner.post('/admin/archive', data={'action': 'run'}, follow_redirects=True)   # 再跑一次：沒有資料
    assert '沒有符合的資料' in r.get_data(as_text=True)
    assert login_staff(app, 'archive_mgr', 'manager', tenant_id=tid).get('/admin/archive').status_code == 403


def test_migration_revision_ids_fit_postgres_alembic_version():
    """alembic_version.version_num 在 PostgreSQL 是 varchar(32)：版本代號不可超過 32 字"""
    import os
    import re as _re
    folder = os.path.join(os.path.dirname(__file__), '..', 'migrations', 'versions')
    for name in os.listdir(folder):
        if name.endswith('.py'):
            src = open(os.path.join(folder, name), encoding='utf-8').read()
            rev = _re.search(r"^revision = '([^']+)'", src, _re.M).group(1)
            assert len(rev) <= 32, (name, rev)


# ===========================================================================
# 多行程與安全細節
# ===========================================================================
def test_scheduler_runs_in_one_process_only(app):
    from app.services import maintenance_service
    now = datetime.now()
    with app.app_context():
        assert maintenance_service.hold_leadership(now, token='proc-a') is True
        assert maintenance_service.hold_leadership(now, token='proc-b') is False      # 另一個 worker
        assert maintenance_service.hold_leadership(now + timedelta(seconds=10), token='proc-a') is True   # 續約
        later = now + timedelta(seconds=maintenance_service.LEADER_LEASE_SECONDS + 15)
        assert maintenance_service.hold_leadership(later, token='proc-b') is True     # A 沒續約 → B 接手
        assert maintenance_service.hold_leadership(later, token='proc-a') is False


def test_campaign_and_webhook_delivery_are_claimed_once(app, monkeypatch):
    from app.models.marketing import MarketingCampaign
    from app.services import marketing_service
    from tests.support import run_together
    tid = _tenant(app, 'claim-shop')
    with _ctx(app, tid):
        c = MarketingCampaign(tenant_id=tid, name='排程中', kind='broadcast', segment='all', message='hi',
                              send_at=datetime.now() - timedelta(minutes=1), status='running')
        db.session.add(c)
        db.session.commit()
        cid = c.id
        with pytest.raises(marketing_service.CampaignBusy):
            marketing_service.run_campaign(c)
        assert marketing_service.run_due(tid) == 0                  # 執行中不重複執行
        db.session.execute(sa_update(MarketingCampaign).where(MarketingCampaign.id == cid)
                           .values(updated_at=datetime.now() - timedelta(hours=1)))
        db.session.commit()
        assert marketing_service.run_due(tid) == 1                  # 卡住超過 30 分鐘：恢復後發送
        assert db.session.get(MarketingCampaign, cid).status == 'done'
        ep = WebhookEndpoint(tenant_id=tid, name='同時送出', url='https://hooks.example.test/x', events='order.created')
        db.session.add(ep)
        db.session.commit()
        for _ in range(5):
            webhook_service.queue_ping(ep)
        db.session.commit()
    monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host: {'93.184.216.34'})
    posted = []

    def worker():
        with app.app_context():
            return webhook_service.deliver_due(tid, post=lambda u, b, h: posted.append(json.loads(b)['id']) or 200)
    results = run_together(worker, worker, worker)
    assert sum(results) == 5 and len(posted) == 5 and len(set(posted)) == 5, (results, posted)


def test_webhook_target_checks_and_pinned_connection(app, monkeypatch):
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer
    got = {}

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            got['host'] = self.headers['Host']
            got['body'] = self.rfile.read(int(self.headers['Content-Length'])).decode()
            got['sig'] = self.headers['X-Kiosk-Signature']
            self.send_response(204)
            self.end_headers()

        def log_message(self, *args):
            pass
    server = HTTPServer(('127.0.0.1', 0), Handler)
    threading.Thread(target=server.handle_request, daemon=True).start()
    port = server.server_address[1]
    with app.app_context():
        for ip in ('10.0.0.5', '127.0.0.1', '169.254.169.254', '100.100.100.200', '::1', '192.168.1.2'):
            monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host, ip=ip: {ip})
            with pytest.raises(Exception, match='不可送到內網'):
                webhook_service.check_public('https://hook.example.test/x')
        monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host: {'93.184.216.34', '10.0.0.1'})
        with pytest.raises(Exception, match='不可送到內網'):     # 其中一個 IP 是內網也拒絕
            webhook_service.check_public('https://hook.example.test/x')
        monkeypatch.setattr(webhook_service, 'resolve_ips', lambda host: {'93.184.216.34'})
        assert webhook_service.check_public('https://hook.example.test/x') == '93.184.216.34'
        # 連線固定到檢查過的 IP，Host 標頭仍是原本的主機名稱 (DNS 之後被換成內網也無效)
        code = webhook_service._post(f'http://hook.example.test:{port}/x?y=1', '{"a": 1}',
                                     {'X-Kiosk-Signature': 'sha256=abc'}, ip='127.0.0.1')
    server.server_close()
    assert code == 204 and got['host'] == f'hook.example.test:{port}' and got['body'] == '{"a": 1}'


def test_webhook_line_pay_authorizing_order_events(app, monkeypatch):
    tid = _tenant(app, 'hook-shop2')
    with _ctx(app, tid):
        ep = WebhookEndpoint(tenant_id=tid, name='ERP', url='https://erp.example.test/h',
                             events=','.join(webhook_service.WEBHOOK_EVENTS), created_at=datetime.now() - timedelta(days=1))
        db.session.add(ep)
        failed = Order(tenant_id=tid, total_price=50, status='Pending', payment_status='authorizing',
                       payment_method='Line Pay', order_type='內用', pickup_number=1)
        ok = Order(tenant_id=tid, total_price=60, status='Pending', payment_status='authorizing',
                   payment_method='Line Pay', order_type='內用', pickup_number=2)
        db.session.add_all([failed, ok])
        db.session.commit()
        assert webhook_service.scan(tid) == 0                      # 扣款中：還沒成立
        failed.status = 'Cancelled'                                # LINE Pay 扣款失敗 → 系統取消
        ok.payment_status = ''                                     # 扣款成功
        db.session.commit()
        webhook_service.scan(tid)
        events = [(json.loads(d.payload)['data']['id'], d.event) for d in
                  WebhookDelivery.query.filter_by(endpoint_id=ep.id).order_by(WebhookDelivery.id).all()]
        ok_id = ok.id
    assert events == [(ok_id, 'order.created'), (ok_id, 'order.paid')], events


def test_exports_neutralize_formulas_and_api_key_not_in_cookie(app):
    from openpyxl import load_workbook
    from app.services import accounting_service
    from app.utils import csv_safe
    assert csv_safe('=HYPERLINK("http://x","y")').startswith("'=") and csv_safe('王小明') == '王小明' and csv_safe(5) == 5
    buf = accounting_service.to_xlsx(['名稱'], [['=1+1'], ['@SUM(A1)']], '測試')
    ws = load_workbook(buf).active
    assert [c.data_type for c in ws['A'][1:]] == ['s', 's'] and ws['A2'].value == '=1+1'
    text = accounting_service.to_csv(['名稱'], [['=1+1']]).read().decode('utf-8')
    assert "'=1+1" in text
    tid = _tenant(app, 'cookie-shop')
    owner = login_staff(app, 'cookie_owner', 'owner', tenant_id=tid)
    r = owner.post('/admin/integrations', data={'action': 'create_key', 'name': 'x', 'scopes': ['menu:read']})
    raw = re.search(r'kk_[A-Za-z0-9_-]+', r.get_data(as_text=True)).group(0)
    assert r.status_code == 200 and r.headers.get('Cache-Control') == 'no-store'
    assert raw not in (r.headers.get('Set-Cookie') or '')
    with owner.session_transaction() as sess:
        assert raw not in json.dumps(dict(sess), default=str)
    assert raw not in owner.get('/admin/integrations').get_data(as_text=True)


def test_archive_keys_rows_by_uuid_even_if_ids_are_reused(app, tmp_path, monkeypatch):
    """SQLite 刪除最大 id 後會重複使用該 id：封存檔以 uuid 為主鍵，不會覆蓋先前封存的資料"""
    import sqlalchemy as sa
    from app.services import archive_service
    tid = _tenant(app, 'archive-reuse')
    monkeypatch.setattr(archive_service, 'archive_dir', lambda: str(tmp_path))
    old = datetime.now() - timedelta(days=3 * 365)
    with _ctx(app, tid):
        engine, meta = archive_service._archive_engine(tid)
        with engine.begin() as conn:
            rows = [type('R', (), {'_mapping': {'id': 99999, 'uuid': f'u-{i}', 'tenant_id': tid, 'item_name': f'品項{i}',
                                                'order_id': 99999, 'price': 1, 'quantity': 1}})() for i in (1, 2)]
            archive_service._copy(conn, meta, 'order_item', rows, {99999: 'o-1'})
            n = conn.execute(sa.text('SELECT COUNT(*) FROM order_item WHERE id = 99999')).scalar()
        engine.dispose()
    assert n == 2
