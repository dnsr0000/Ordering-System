"""第四輪：營業費用登錄、損益表、原料成本 / 售價歷史、對外 API 說明文件"""
import io
import re
from contextlib import contextmanager
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models.expense import CostHistory, Expense
from app.models.inventory import Ingredient, RecipeLine
from app.models.menu import MenuItem
from app.models.payment import PaymentRecord
from app.models.tenant import StaffAccount, Tenant
from app.services import pnl_service
from tests.support import assert_flash, login_staff


def _tenant(app, slug):
    with app.app_context():
        t = Tenant.query.filter_by(slug=slug).first()
        if t is None:
            t = Tenant(name=f'測試 {slug}', slug=slug)
            db.session.add(t)
            db.session.commit()
        return t.id


@contextmanager
def _ctx(app, tenant_id):
    from app.tenant_context import _use_tenant
    with app.test_request_context('/'):
        _use_tenant(db.session.get(Tenant, tenant_id), remember=False)
        yield


def _expenses(app, tid, **filters):
    with app.app_context():
        return [(e.name, e.category, e.amount, e.spent_on, e.is_fixed, e.is_deleted)
                for e in Expense.query.filter_by(tenant_id=tid, **filters).order_by(Expense.id)]


# ===========================================================================
# 營業費用
# ===========================================================================
def test_expense_add_edit_delete_copy_fixed_and_permissions(app):
    tid = _tenant(app, 'exp-shop')
    owner = login_staff(app, 'exp_owner', role='owner', tenant_id=tid)
    base = dict(action='add', spent_on='2026-08-05', category='rent', name='八月房租', amount='30000', vendor='房東',
                invoice_no='ab12345678', is_fixed='1')
    r = owner.post('/admin/expenses?month=2026-08', data=base, follow_redirects=True)
    assert r.status_code == 200 and '已新增費用「八月房租」$30,000' in r.get_data(as_text=True)
    owner.post('/admin/expenses?month=2026-08', data=dict(base, category='utilities', name='電費', amount='4200',
                                                         is_fixed=''))
    rows = _expenses(app, tid)
    assert rows == [('八月房租', 'rent', 30000, date(2026, 8, 5), True, False),
                    ('電費', 'utilities', 4200, date(2026, 8, 5), False, False)]
    with app.app_context():
        rent = Expense.query.filter_by(tenant_id=tid, name='八月房租').one()
        assert rent.invoice_no == 'AB12345678' and rent.created_by == 'exp_owner' and rent.node_uuid is None
        rent_id, power_id = rent.id, Expense.query.filter_by(tenant_id=tid, name='電費').one().id

    # 驗證：金額、類別、名稱、日期
    for bad, msg in ((dict(amount='0'), '正確的金額'), (dict(amount='-5'), '正確的金額'), (dict(category='xx'), '費用類別'),
                     (dict(name=' '), '費用名稱'), (dict(spent_on='2026/13/40'), '請輸入日期')):
        owner.post('/admin/expenses', data=dict(base, **bad))
        assert_flash(owner, msg)
    assert len(_expenses(app, tid)) == 2

    # 修改、刪除 (標記)
    owner.post('/admin/expenses?month=2026-08', data=dict(base, action='edit', expense_id=str(power_id), name='電費',
                                                         category='utilities', amount='4500', is_fixed=''))
    owner.post('/admin/expenses?month=2026-08', data={'action': 'delete', 'expense_id': str(power_id)})
    assert_flash(owner, '已刪除「電費」')
    assert ('電費', 'utilities', 4500, date(2026, 8, 5), False, True) in _expenses(app, tid)
    page = owner.get('/admin/expenses?month=2026-08').get_data(as_text=True)
    assert '八月房租' in page and '電費' not in page and '30,000' in page

    # 下個月一鍵複製固定費用 (只複製勾選「每月固定」的，重複按不會再複製)
    owner.post('/admin/expenses?month=2026-09', data={'action': 'copy_fixed'})
    assert ('八月房租', 'rent', 30000, date(2026, 9, 5), True, False) in _expenses(app, tid)
    owner.post('/admin/expenses?month=2026-09', data={'action': 'copy_fixed'})
    assert_flash(owner, '沒有尚未複製的固定費用')
    assert len(_expenses(app, tid, spent_on=date(2026, 9, 5))) == 1
    # 1/31 的固定費用複製到 2 月 → 2/28
    with app.app_context():
        db.session.add(Expense(tenant_id=tid, spent_on=date(2027, 1, 31), category='other', name='月底清潔', amount=800,
                               is_fixed=True))
        db.session.commit()
    owner.post('/admin/expenses?month=2027-02', data={'action': 'copy_fixed'})
    assert ('月底清潔', 'other', 800, date(2027, 2, 28), True, False) in _expenses(app, tid)

    # 其他店家的費用：查無此費用；一般員工不可進入；未知操作
    other = _tenant(app, 'exp-other')
    other_owner = login_staff(app, 'exp_other_owner', role='owner', tenant_id=other)
    assert other_owner.post('/admin/expenses', data={'action': 'delete', 'expense_id': str(rent_id)}).status_code == 404
    assert not _expenses(app, tid, id=rent_id, is_deleted=True)
    staff = login_staff(app, 'exp_staff', role='staff', tenant_id=tid)
    assert staff.get('/admin/expenses').status_code in (302, 403)
    owner.post('/admin/expenses', data={'action': 'hack'})
    assert_flash(owner, '不支援的操作')


# ===========================================================================
# 損益表
# ===========================================================================
def test_pnl_statement_numbers_fees_tax_trend_and_excel(app):
    from app.services import report_service, settings_service
    from app.services.order_service import process_checkout
    tid = _tenant(app, 'pnl-shop')
    today = date.today()
    month = f'{today:%Y-%m}'
    start, end = pnl_service.month_range(month)
    now = datetime.now()
    with _ctx(app, tid):
        item = MenuItem(tenant_id=tid, name='損益牛排', price=1050, unit_cost=300, stock=99, total_stock=99)
        db.session.add(item)
        db.session.commit()
        process_checkout(tenant_id=tid, items=[{'menu_item_id': item.id, 'quantity': 2}])
        db.session.commit()
        for kind, method, amount in (('sale', 'Credit Card', 1000), ('settle', 'Line Pay', 2000),
                                     ('sale', '平台代收', 500), ('topup', 'Apple Pay', 300), ('refund', 'Credit Card', -400),
                                     ('pay_out', 'Cash', -100)):
            db.session.add(PaymentRecord(tenant_id=tid, kind=kind, method=method, amount=amount, created_at=now))
        db.session.add_all([
            Expense(tenant_id=tid, spent_on=start, category='rent', name='房租', amount=5000),
            Expense(tenant_id=tid, spent_on=start, category='utilities', name='水電', amount=1200),
            Expense(tenant_id=tid, spent_on=start, category='other', name='已刪除', amount=999, is_deleted=True),
            Expense(tenant_id=tid, spent_on=start - timedelta(days=1), category='rent', name='上月房租', amount=4000),
        ])
        db.session.commit()
        r = report_service.build_report(tid, start, end)
        assert r['net_after_refunds'] == 2100 and r['cogs'] == 600
        s = pnl_service.statement(tid, start, end)
        assert (s['gross_sales'], s['tax'], s['revenue'], s['cogs'], s['gross_profit']) == (2100, 100, 2000, 600, 1400)
        # 刷卡 2% × (1000 + 300)、行動支付 2.5% × 2000、外送平台 30% × 500；退款 / 支出不計
        assert s['fees'] == {'card_fee_percent': 26, 'mobile_fee_percent': 50, 'online_fee_percent': 0,
                             'platform_fee_percent': 150}
        assert s['fee_total'] == 226 and s['expenses']['rent'] == 5000 and s['expense_total'] == 6200
        assert s['labor'] == 0 and s['opex'] == 6426 and s['net'] == 1400 - 6426
        assert s['gross_margin'] == 70.0 and s['cogs_ratio'] == 30.0 and s['net_margin'] == round(-5026 / 2000 * 100, 1)
        # 費率可在系統設定調整；稅率 0 時營業額即收入
        settings_service.put('pnl', {'card_fee_percent': 3.0, 'mobile_fee_percent': 0.0, 'online_fee_percent': 0.0,
                                     'platform_fee_percent': 0.0}, tenant_id=tid)
        settings_service.put('accounting', dict(settings_service.get('accounting', tid), tax_rate=0), tenant_id=tid)
        s2 = pnl_service.statement(tid, start, end)
        assert s2['fee_total'] == 39 and s2['tax'] == 0 and s2['revenue'] == 2100
        settings_service.put('accounting', dict(settings_service.get('accounting', tid), tax_rate=5), tenant_id=tid)
        settings_service.put('pnl', {'card_fee_percent': 2.0, 'mobile_fee_percent': 2.5, 'online_fee_percent': 2.75,
                                     'platform_fee_percent': 30.0}, tenant_id=tid)
        db.session.commit()
        prev = pnl_service.statement(tid, *pnl_service.previous_period(start, end))
        assert prev['expenses']['rent'] == 4000 and prev['net'] == -4000
        months = pnl_service.trend(tid, 6)
        assert len(months) == 6 and months[-1]['month'] == month and months[-1]['net'] == s['net']
        assert months[-2]['net'] == -4000

    assert pnl_service.previous_period(date(2026, 10, 1), date(2026, 10, 31)) == (date(2026, 9, 1), date(2026, 9, 30))
    assert pnl_service.previous_period(date(2026, 3, 1), date(2026, 3, 31)) == (date(2026, 2, 1), date(2026, 2, 28))
    assert pnl_service.previous_period(date(2026, 10, 5), date(2026, 10, 11)) == (date(2026, 9, 28), date(2026, 10, 4))
    assert pnl_service.month_range('亂打') == pnl_service.month_range(f'{today:%Y-%m}')
    # 跨多個整月：往前同樣月數 (一季比上一季、一年比去年)
    assert pnl_service.previous_period(date(2026, 7, 1), date(2026, 9, 30)) == (date(2026, 4, 1), date(2026, 6, 30))
    assert pnl_service.previous_period(date(2026, 1, 1), date(2026, 12, 31)) == (date(2025, 1, 1), date(2025, 12, 31))
    # 極端年份不會溢位
    assert pnl_service.month_range('9999-12') == pnl_service.month_range(f'{today:%Y-%m}')

    owner = login_staff(app, 'pnl_owner', role='owner', tenant_id=tid)
    page = owner.get(f'/admin/pnl?month={month}').get_data(as_text=True)
    assert '營業淨利' in page and '-$5,026' in page and '刷卡手續費' in page and '上月房租' not in page
    custom = owner.get(f'/admin/pnl?start={end}&end={start}')   # 起訖顛倒自動交換
    assert custom.status_code == 200 and '-$5,026' in custom.get_data(as_text=True)
    owner.get('/admin/pnl?start=2020-01-01&end=2026-01-01')
    assert_flash(owner, '最長 400 天')
    owner.get('/admin/pnl?start=2026-13-01&end=2026-01-01')
    assert_flash(owner, '日期格式錯誤')
    owner.get('/admin/pnl?start=0001-01-01&end=0001-01-31')
    assert_flash(owner, '2000 ~ 2100')
    assert owner.get('/admin/pnl?month=9999-12').status_code == 200
    assert 'pnlNoLabor' not in page   # 單機有打卡紀錄：人事成本列入
    x = owner.get(f'/admin/pnl?month={month}&format=xlsx')
    assert x.status_code == 200 and 'spreadsheetml' in x.mimetype
    wb = load_workbook(io.BytesIO(x.data))
    rows = {row[0]: row[1:] for row in wb['損益表'].iter_rows(values_only=True) if row and row[0]}
    assert rows['營業淨利'][0] == -5026 and rows['  刷卡手續費'][0] == 26 and rows['  租金'][:2] == (5000, 4000)
    assert [r[0] for r in wb['近六個月'].iter_rows(min_row=2, values_only=True)][-1] == month
    # 沒有報表權限的員工不可看損益表
    kitchen = login_staff(app, 'pnl_kitchen', role='kitchen', tenant_id=tid)
    assert kitchen.get('/admin/pnl').status_code in (302, 403)


def test_pnl_excel_neutralizes_formula_text(app):
    cur = {k: 0 for _l, k in pnl_service.STATEMENT_ROWS}
    cur.update(start=date(2026, 1, 1), end=date(2026, 1, 31), net_margin=0.0,
               fees={k: 0 for k in pnl_service.FEE_LABELS}, expenses={k: 0 for k in pnl_service.EXPENSE_CATEGORIES})
    wb = load_workbook(pnl_service.to_excel(cur, cur, [], '=HYPERLINK("http://x")'))
    cell = wb['損益表']['B1']
    assert cell.value == '=HYPERLINK("http://x")' and cell.data_type == 's'   # 以文字儲存，Excel 不會當成公式


# ===========================================================================
# 原料成本 / 售價歷史
# ===========================================================================
def test_cost_history_from_purchase_edit_and_price_change(app):
    from flask_login import login_user
    from app.services import inventory_service
    tid = _tenant(app, 'cost-shop')
    login_staff(app, 'cost_owner', role='owner', tenant_id=tid)
    with _ctx(app, tid):
        login_user(StaffAccount.query.filter_by(username='cost_owner').one())
        beef = Ingredient(tenant_id=tid, name='成本牛肉', unit='kg', stock=100, unit_cost=200)
        noodle = MenuItem(tenant_id=tid, name='成本牛肉麵', price=150, stock=10, total_stock=10)
        db.session.add_all([beef, noodle])
        db.session.commit()
        db.session.add(RecipeLine(tenant_id=tid, menu_item_id=noodle.id, ingredient_id=beef.id, quantity=0.2))
        db.session.commit()
        assert CostHistory.query.filter_by(tenant_id=tid).count() == 0   # 新增不算變動
        # 進貨加權平均：(200×100 + 260×100) / 200 = 230
        inventory_service.adjust(beef, 100, 'purchase', unit_cost=260)
        db.session.commit()
        noodle.price = 160
        noodle.name = noodle.name   # 未變動的欄位不記錄
        db.session.commit()
        noodle.price = 160
        db.session.commit()
        rows = [(h.kind, h.field, h.old_value, h.new_value, h.changed_by)
                for h in CostHistory.query.filter_by(tenant_id=tid).order_by(CostHistory.id)]
        assert rows == [('ingredient', 'unit_cost', 200, 230, 'cost_owner'), ('menu', 'price', 150, 160, 'cost_owner')]
        assert CostHistory.query.filter_by(tenant_id=tid, kind='menu').one().change_pct == 6.7
        # 未載入舊值時 (直接以 id 取得新物件) 以資料庫的值為準
        db.session.expire_all()
        db.session.get(Ingredient, beef.id).unit_cost = 250
        db.session.commit()
        assert CostHistory.query.filter_by(tenant_id=tid, ref_id=beef.id).order_by(CostHistory.id.desc()).first() \
            .old_value == 230
        # 成本變化對毛利的影響：與 1 天前 (200/kg) 相比，現在 250/kg → 每份 40 → 50
        then = datetime.now() - timedelta(days=1)
        assert pnl_service.ingredient_cost_at(tid, then)[beef.id] == 200
        impact = pnl_service.cost_impact(tid, days=1)
        assert impact == [{'id': noodle.id, 'name': '成本牛肉麵', 'price': 160, 'cost_then': 40.0, 'cost_now': 50.0,
                           'delta': 10.0, 'margin_then': 75.0, 'margin_now': 68.8}]
        # 一年前的變動不在 30 天內的比較
        CostHistory.query.filter_by(tenant_id=tid).update({'changed_at': datetime.now() - timedelta(days=400)})
        db.session.commit()
        assert pnl_service.cost_impact(tid, days=30) == []
        CostHistory.query.filter_by(tenant_id=tid).update({'changed_at': datetime.now()})
        db.session.commit()

    owner = login_staff(app, 'cost_owner', role='owner', tenant_id=tid)
    page = owner.get('/admin/cost-history?days=30').get_data(as_text=True)
    assert '成本牛肉麵' in page and '成本牛肉' in page and '單位成本' in page and '售價' in page
    only_menu = owner.get('/admin/cost-history?kind=menu').get_data(as_text=True)
    assert 'historyTable' in only_menu and '單位成本' not in only_menu.split('historyTable')[1]
    assert owner.get('/admin/cost-history?days=abc').status_code == 200


def test_settings_page_lists_round4_pages(admin_client):
    html = admin_client.get('/admin/settings').get_data(as_text=True)
    for name in ('損益表', '營業費用', '成本 / 售價歷史'):
        assert name in html


# ===========================================================================
# 多分店：分店登錄的費用上傳到線上端，損益表可依分店篩選
# ===========================================================================
@pytest.fixture(scope='module')
def env():
    from tests.test_sync_hub_node import Env
    e = Env()
    yield e
    e.close()


def test_node_expenses_upload_to_hub_and_filter_by_branch(env):
    from tests.test_sync_hub_node import tenant_ctx
    month = f'{date.today():%Y-%m}'
    day = date.today().replace(day=1).isoformat()
    a_owner = login_staff(env.a, 'a_exp_owner', role='owner')
    a_owner.post('/admin/expenses', follow_redirects=True, data={'action': 'add', 'spent_on': day, 'category': 'rent', 'name': '台南店房租',
                                         'amount': '25000'})
    b_owner = login_staff(env.b, 'b_exp_owner', role='owner')
    b_owner.post('/admin/expenses', follow_redirects=True, data={'action': 'add', 'spent_on': day, 'category': 'utilities', 'name': '高雄店電費',
                                         'amount': '3000'})
    hub_owner = login_staff(env.hub, 'hub_exp_owner', role='owner')
    hub_owner.post('/admin/expenses', follow_redirects=True, data={'action': 'add', 'spent_on': day, 'category': 'marketing', 'name': '總部廣告',
                                           'amount': '8000'})
    env.sync_all()
    with tenant_ctx(env.hub):
        got = {e.name: e.node_uuid for e in Expense.query.all()}
        assert got == {'台南店房租': env.node_a_uuid, '高雄店電費': env.node_b_uuid, '總部廣告': None}
        a_id = Expense.query.filter_by(name='台南店房租').one().id
    with tenant_ctx(env.a):
        assert [e.name for e in Expense.query.all()] == ['台南店房租']   # 總部 / 其他分店的費用不會下傳
    page = hub_owner.get(f'/admin/pnl?month={month}&node={env.node_a_uuid}').get_data(as_text=True)
    assert '25,000' in page and '8,000' not in page and 'pnlNoLabor' in page   # 線上端：人事成本未列入的提示
    whole = hub_owner.get(f'/admin/pnl?month={month}').get_data(as_text=True)
    assert '36,000' in whole
    listing = hub_owner.get(f'/admin/expenses?month={month}').get_data(as_text=True)
    assert '台南店' in listing and '高雄店' in listing
    # 線上端「複製上個月固定費用」只複製線上端自己的 (分店的房租由分店自己複製後上傳，不會重複計算)
    from app.models.expense import Expense as Exp
    with tenant_ctx(env.hub) as tenant:
        prev = (date.today().replace(day=1) - timedelta(days=1)).replace(day=1)
        for name, node in (('上月分店房租', env.node_a_uuid), ('上月總部倉租', None)):
            db.session.add(Exp(tenant_id=tenant.id, spent_on=prev, category='rent', name=name, amount=1000, is_fixed=True,
                               node_uuid=node))
        db.session.commit()
        assert pnl_service.copy_fixed(tenant.id, month) == 1
        db.session.rollback()
    # 分店不可藉同步改別家分店 (或線上端) 的費用：用相同 uuid 上傳會被拒絕
    with tenant_ctx(env.hub):
        b_uuid = Expense.query.filter_by(name='高雄店電費').one().uuid
    with tenant_ctx(env.a) as tenant:
        db.session.add(Expense(tenant_id=tenant.id, uuid=b_uuid, spent_on=date.today(), category='utilities',
                               name='偷改的電費', amount=1, node_uuid=env.node_a_uuid))
        db.session.commit()
    env.sync_all()
    with tenant_ctx(env.hub):
        hacked = Expense.query.filter_by(uuid=b_uuid).one()
        assert (hacked.name, hacked.amount, hacked.node_uuid) == ('高雄店電費', 3000, env.node_b_uuid)
    with tenant_ctx(env.a):
        Expense.query.filter_by(uuid=b_uuid).delete()
        db.session.commit()
    # 線上端不可改分店登錄的費用；分店刪除後同步到線上端
    assert hub_owner.post('/admin/expenses', data={'action': 'delete', 'expense_id': str(a_id)}).status_code == 403
    with tenant_ctx(env.a):
        local_id = Expense.query.one().id
    a_owner.post('/admin/expenses', data={'action': 'delete', 'expense_id': str(local_id)}, follow_redirects=True)
    env.sync_all()
    with tenant_ctx(env.hub):
        assert Expense.query.filter_by(name='台南店房租').one().is_deleted is True
    assert '25,000' not in hub_owner.get(f'/admin/pnl?month={month}&node={env.node_a_uuid}').get_data(as_text=True)


# ===========================================================================
# 對外 API 說明文件
# ===========================================================================
def test_developer_docs_page_and_openapi_cover_every_api_route(app):
    from app.models.integration import API_SCOPES, WEBHOOK_EVENTS
    client = app.test_client()
    page = client.get('/developers')
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    for text in ('/api/v1/orders', 'X-Kiosk-Signature', 'Authorization: Bearer', 'tryItForm', 'hmac.compare_digest',
                 *API_SCOPES, *WEBHOOK_EVENTS):
        assert text in html
    assert '<script>' not in html.split('id="page-data">')[1].split('</script>')[0]

    spec = client.get('/developers/openapi.json')
    assert spec.status_code == 200 and spec.headers['Access-Control-Allow-Origin'] == '*'
    doc = spec.get_json()
    assert doc['openapi'].startswith('3.') and doc['components']['securitySchemes']['bearerAuth']['scheme'] == 'bearer'
    documented = {(m.upper(), p) for p, ops in doc['paths'].items() for m in ops}
    actual = set()
    for rule in app.url_map.iter_rules():
        if rule.endpoint.startswith('api_v1.'):
            path = re.sub(r'<(?:\w+:)?(\w+)>', r'{\1}', rule.rule)
            actual |= {(m, path) for m in rule.methods - {'HEAD', 'OPTIONS'}}
    assert documented == actual
    # 每個 $ref 都指得到
    text = spec.get_data(as_text=True)
    for ref in set(re.findall(r'#/components/(schemas|responses)/(\w+)', text)):
        assert ref[1] in doc['components'][ref[0]]
    ops = [op for ops in doc['paths'].values() for op in ops.values()]
    assert len({op['operationId'] for op in ops}) == len(ops)
    assert {p['name'] for p in doc['paths']['/api/v1/orders']['get']['parameters']} >= {'since', 'after_id', 'cursor'}
    assert set(doc['x-webhooks']) == set(WEBHOOK_EVENTS)


def test_integrations_page_links_to_docs(app):
    owner = login_staff(app, 'docs_owner', role='owner')
    assert '/developers' in owner.get('/admin/integrations').get_data(as_text=True)
